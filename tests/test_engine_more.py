"""More engine behaviour: structural edits keep every reference right, sorting and filtering,
conditional-format evaluation, tables, names, protection and XLSX round trips of features."""

import base64
import io
import time
import warnings

import openpyxl
import pytest

from gridcore.core.addr import RangeRef
from gridcore.engine import EngineError, Session, objects, viewport
from gridcore.engine.cf import CfCache
from gridcore.model.features import CHART_KINDS


def new():
    s = Session()
    s.new_workbook()
    return s


def val(s, a):
    return s.run("cell.get", cell=a)["value"]


def inp(s, a):
    return s.run("cell.get", cell=a)["input"]


def roundtrip(s) -> Session:
    data = s.run("file.saveBytes")["data"]
    s2 = Session()
    s2.run("file.open", data=data, name="rt.xlsx")
    return s2


def openpyxl_book(s, data_only=False):
    raw = base64.b64decode(s.run("file.saveBytes")["data"])
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(io.BytesIO(raw), data_only=data_only)


# ---------------------------------------------------------------- structure


def test_insert_columns_moves_everything():
    s = new()
    s.run("range.setValues", range="A1", values=[[1, 2, 3]])
    s.run("cell.set", cell="D1", input="=SUM(A1:C1)")
    s.run("home.mergeCenter", range="B3:C3")
    s.run("review.newNote", cell="C1", text="n")
    s.run("home.conditionalFormat", range="C1", rule={"type": "cellIs", "operator": "greaterThan", "value": "0"})
    s.run("formulas.defineName", name="Third", refersTo="=Sheet1!$C$1")
    s.run("home.insertColumns", cols="B")
    sh = s.doc().sheet
    assert inp(s, "E1") == "=SUM(A1:D1)" and val(s, "E1") == 6
    assert sh.merges == [RangeRef(2, 2, 2, 3)]
    assert (0, 3) in sh.comments
    assert sh.cond_formats[0].ranges == [RangeRef(0, 3, 0, 3)]
    assert s.doc().wb.names[0].formula == "Sheet1!$D$1"


def test_delete_rows_shrinks_ranges_and_refs():
    s = new()
    s.run("range.setValues", range="A1", values=[[1], [2], [3], [4]])
    s.run("cell.set", cell="B1", input="=SUM(A1:A4)")
    s.run("home.deleteRows", rows="2:3")
    assert inp(s, "B1") == "=SUM(A1:A2)" and val(s, "B1") == 5


def test_insert_and_delete_cells_shift():
    s = new()
    s.run("range.setValues", range="A1", values=[[1], [2], [3]])
    s.run("cell.set", cell="C1", input="=A3")
    s.run("selection.set", range="A1")
    s.run("home.insertCells", shift="down")
    assert val(s, "A2") == 1 and inp(s, "C1") == "=A4"
    s.run("home.deleteCells", shift="up", range="A1")
    assert val(s, "A1") == 1 and inp(s, "C1") == "=A3"


def test_move_cells_keeps_references():
    s = new()
    s.run("range.setValues", range="A1", values=[[5, "=A1*2"]])
    s.run("cell.set", cell="C1", input="=B1+1")
    s.run("edit.moveCells", **{"from": "A1:B1", "to": "A5"})
    assert inp(s, "B5") == "=A5*2"
    assert inp(s, "C1") == "=B5+1" and val(s, "C1") == 11


def test_rename_move_copy_sheets():
    s = new()
    s.run("cell.set", cell="A1", input="3")
    s.run("home.insertSheet", name="Calc", index=1)
    s.run("cell.set", cell="A1", input="=Sheet1!A1*2")
    s.run("sheet.duplicate")
    names = [sh.name for sh in s.doc().wb.sheets]
    assert names == ["Sheet1", "Calc", "Calc (2)"]
    s.run("sheet.move", sheet="Calc (2)", to=0)
    assert s.doc().wb.sheets[0].name == "Calc (2)"
    assert val(s, "A1") == 6
    with pytest.raises(EngineError):
        s.run("sheet.rename", sheet="Calc", name="Bad/Name")


def test_undo_restores_structure():
    s = new()
    s.run("range.setValues", range="A1", values=[[1], [2]])
    s.run("home.deleteRows", rows="1")
    assert val(s, "A1") == 2
    s.run("edit.undo")
    assert val(s, "A1") == 1 and val(s, "A2") == 2


# ---------------------------------------------------------------- data


def test_sort_keeps_rows_together_and_moves_formulas():
    s = new()
    s.run("range.setValues", range="A1", values=[["k", "v", "w"], ["b", 2, "=B2*10"], ["a", 1, "=B3*10"], ["c", 3, "=B4*10"]])
    s.run("data.sort", range="A1:C4", keys=[{"col": "A"}], header=True)
    assert [val(s, f"A{i}") for i in (2, 3, 4)] == ["a", "b", "c"]
    assert [val(s, f"C{i}") for i in (2, 3, 4)] == [10, 20, 30]
    s.run("data.sort", range="A1:C4", keys=[{"col": "B", "order": "desc"}], header=True)
    assert [val(s, f"B{i}") for i in (2, 3, 4)] == [3, 2, 1]


def test_sort_puts_blanks_last_both_ways():
    s = new()
    s.run("range.setValues", range="A1", values=[[3], [None], [1], ["x"]])
    s.run("data.sort", range="A1:A4", keys=[{"col": "A"}], header=False)
    assert [val(s, f"A{i}") for i in range(1, 5)] == [1, 3, "x", None]
    s.run("data.sort", range="A1:A4", keys=[{"col": "A", "order": "desc"}], header=False)
    assert [val(s, f"A{i}") for i in range(1, 5)] == ["x", 3, 1, None]


def test_filter_values_top10_and_clear():
    s = new()
    s.run("range.setValues", range="A1", values=[["n"]] + [[i] for i in range(1, 11)])
    s.run("selection.set", range="A1")
    s.run("data.filter")
    s.run("data.filterColumn", col="A", criterion={"kind": "top10", "count": 3})
    sh = s.doc().sheet
    hidden = [r for r in range(1, 11) if sh.is_row_hidden(r)]
    assert hidden == list(range(1, 8))
    s.run("data.filterColumn", col="A", criterion={"kind": "values", "values": ["2", "4"]})
    sh = s.doc().sheet
    assert [r + 1 for r in range(1, 11) if not sh.is_row_hidden(r)] == [3, 5]
    s.run("data.clearFilter")
    assert not any(s.doc().sheet.is_row_hidden(r) for r in range(11))


def test_remove_duplicates():
    s = new()
    s.run("range.setValues", range="A1", values=[["k", "v"], ["a", 1], ["A", 1], ["b", 2], ["a", 3]])
    r = s.run("data.removeDuplicates", range="A1:B5", columns=[0], header=True)
    assert r == {"removed": 2, "remaining": 2}
    assert [val(s, f"A{i}") for i in (2, 3, 4)] == ["a", "b", None]


def test_text_to_columns():
    s = new()
    s.run("range.setValues", range="A1", values=[["x,1,2026-01-02"], ["y,2,2026-01-03"]])
    s.run("data.textToColumns", range="A1:A2", delimiters=[","])
    assert val(s, "A2") == "y" and val(s, "B2") == 2 and s.run("cell.get", cell="C1")["text"] == "1/2/2026"


def test_validation_kinds():
    s = new()
    s.run("data.validation", range="A1", kind="whole", operator="between", f1="1", f2="10")
    s.run("cell.set", cell="A1", input="5")
    with pytest.raises(EngineError):
        s.run("cell.set", cell="A1", input="5.5")
    s.run("data.validation", range="B1", kind="textLength", operator="lessThanOrEqual", f1="3")
    with pytest.raises(EngineError):
        s.run("cell.set", cell="B1", input="abcd")
    s.run("data.validation", range="C1", kind="custom", f1="=C1>A1")
    s.run("cell.set", cell="C1", input="6")
    with pytest.raises(EngineError):
        s.run("cell.set", cell="C1", input="4")
    s.run("data.validation", range="D1", kind="decimal", operator="greaterThan", f1="0", errorStyle="warning")
    s.run("cell.set", cell="D1", input="-1")  # warnings let the value through
    assert val(s, "D1") == -1


# ---------------------------------------------------------------- conditional formats


def test_cf_rules_evaluate():
    s = new()
    s.run("range.setValues", range="A1", values=[[1], [5], [5], [9]])
    s.run("home.conditionalFormat", range="A1:A4", rule={"type": "duplicateValues", "preset": "yellow"})
    s.run("home.conditionalFormat", range="A1:A4", rule={"type": "top10", "rank": 1, "preset": "green"})
    s.run("home.conditionalFormat", range="A1:A4", rule={"type": "colorScale", "scale": "whiteRed"})
    s.run("home.conditionalFormat", range="A1:A4", rule={"type": "expression", "formula": "=A1<2", "preset": "redText"})
    cf = CfCache(s.doc().wb, 0)
    a1, a2, a4 = cf.look((0, 0)), cf.look((1, 0)), cf.look((3, 0))
    assert a2.style.fill.fg.value == 0xFFEB9C  # duplicate
    assert a4.style.fill.fg.value == 0xC6EFCE  # top 1
    assert a1.style.font.color.value == 0x9C0006  # formula
    assert a1.scale_fill == (0xFC, 0xFC, 0xFF) and a4.scale_fill == (0xF8, 0x69, 0x6B)


def test_cf_icons_and_bars_in_viewport():
    s = new()
    s.run("range.setValues", range="A1", values=[[0], [50], [100]])
    s.run("home.conditionalFormat", range="A1:A3", rule={"type": "iconSet", "iconSet": "3Arrows"})
    s.run("home.conditionalFormat", range="B1:B3", rule={"type": "dataBar"})
    vp = viewport(s.doc().wb, 0, RangeRef(0, 0, 5, 5))
    icons = {c["r"]: c["x"]["icon"][1] for c in vp["cells"] if "x" in c and "icon" in c["x"]}
    assert icons == {0: 0, 1: 1, 2: 2}


# ---------------------------------------------------------------- tables and names


def test_table_calculated_column_and_totals_roundtrip():
    s = new()
    s.run("range.setValues", range="A1", values=[["Qty", "Price", "Total"], [2, 3, None], [4, 5, None]])
    s.run("selection.set", range="A1")
    s.run("insert.table")
    s.run("cell.set", cell="C2", input="=[@Qty]*[@Price]")
    assert val(s, "C3") == 20
    s.run("selection.set", range="A2")
    s.run("table.totalRow")
    assert val(s, "C4") == 26
    s2 = roundtrip(s)
    t = s2.doc().sheet.tables[0]
    assert t.totals_row and t.columns[2].formula is not None
    assert inp(s2, "C2") == "=[@Qty]*[@Price]" and val(s2, "C4") == 26
    wb = openpyxl_book(s)
    assert wb.active["C2"].value == "=Table1[[#This Row],[Qty]]*Table1[[#This Row],[Price]]"


def test_sheet_scoped_name_and_lambda_name():
    s = new()
    s.run("formulas.defineName", name="Rate", refersTo="=0.1", scope="Sheet1")
    s.run("formulas.defineName", name="DOUBLE", refersTo="=LAMBDA(x,x*2)")
    s.run("cell.set", cell="A1", input="=100*Rate")
    s.run("cell.set", cell="A2", input="=DOUBLE(21)")
    assert val(s, "A1") == 10 and val(s, "A2") == 42
    s2 = roundtrip(s)
    assert val(s2, "A1") == 10 and val(s2, "A2") == 42


# ---------------------------------------------------------------- xlsx features


def test_styles_merges_freeze_widths_roundtrip():
    s = new()
    s.run("range.setValues", range="A1", values=[["Head", 1234.567]])
    s.run("home.bold", range="A1")
    s.run("home.fillColor", range="A1", color="#FFC000")
    s.run("home.fontColor", range="A1", color="theme:4:400")
    s.run("home.numberFormat", range="B1", code="#,##0.00")
    s.run("home.borders", range="A1:B1", preset="thickOutside")
    s.run("home.mergeCells", range="C1:D2")
    s.run("home.columnWidth", cols="A", width=150)
    s.run("home.rowHeight", rows="3", height=40)
    s.run("view.freezePanes", cell="B2")
    s.run("home.wrapText", range="A1")
    wb = openpyxl_book(s)
    ws = wb.active
    assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb in ("FFFFC000", "00FFC000")
    assert ws["A1"].alignment.wrap_text
    assert ws["B1"].number_format == "#,##0.00"
    assert ws["A1"].border.left.style == "medium"
    assert "C1:D2" in [str(r) for r in ws.merged_cells.ranges]
    assert ws.freeze_panes == "B2"
    # 150 px = 20.71 characters; the file stores the width with Excel's 5 px padding included.
    assert abs(ws.column_dimensions["A"].width - 21.43) < 0.02
    assert abs(ws.row_dimensions[3].height - 30) < 0.6
    s2 = roundtrip(s)
    st = s2.doc().wb.styles.get(s2.doc().sheet.cell((0, 0)).style)
    assert st.font.bold and st.font.color.kind == "theme" and st.font.color.value == 4
    assert s2.doc().sheet.freeze == (1, 1)


def test_validation_cf_links_notes_charts_roundtrip():
    s = new()
    s.run("range.setValues", range="A1", values=[["", "Q1", "Q2"], ["East", 3, 4], ["West", 5, 6]])
    s.run("data.validation", range="E1:E9", kind="list", f1="Yes,No", errorMessage="Pick one")
    s.run("home.conditionalFormat", range="B2:C3", rule={"type": "colorScale", "scale": "greenYellowRed"})
    s.run("insert.hyperlink", cell="F1", target="https://example.com", text="site")
    s.run("review.newNote", cell="F2", text="Remember")
    s.run("selection.set", range="A1:C3")
    s.run("insert.chart", kind="barStacked", title="Sales")
    s2 = roundtrip(s)
    sh = s2.doc().sheet
    assert sh.validations[0].kind == "list" and sh.validations[0].error_message == "Pick one"
    assert sh.cond_formats[0].rule.type == "colorScale" and len(sh.cond_formats[0].rule.stops) == 3
    assert sh.hyperlinks[(0, 5)].target == "https://example.com"
    assert sh.comments[(1, 5)].text == "Remember"
    ch = objects(s2.doc().wb, 0)["charts"][0]
    assert ch["kind"] == "barStacked" and ch["title"] == "Sales" and [x["name"] for x in ch["series"]] == ["Q1", "Q2"]
    wb = openpyxl_book(s)
    ws = wb.active
    assert ws["F1"].hyperlink.target == "https://example.com"
    assert ws["F2"].comment.text.startswith("Remember")
    assert len(ws._charts) == 1


@pytest.mark.parametrize("kind", CHART_KINDS)
def test_every_chart_kind_roundtrips(kind):
    s = new()
    s.run("range.setValues", range="A1", values=[["", "Q1", "Q2"], ["East", 3, 4], ["West", 5, 6]])
    s.run("selection.set", range="A1:C3")
    s.run("insert.chart", kind=kind)
    assert objects(roundtrip(s).doc().wb, 0)["charts"][0]["kind"] == kind
    ws = openpyxl_book(s).active
    assert len(ws._charts) == 1
    if kind.startswith("scatter"):
        # A scatter chart without lines is one whose series lines have no fill, as Excel writes it.
        assert bool(ws._charts[0].series[0].spPr.ln.noFill) == (kind == "scatter")


def test_a_cell_holds_a_note_or_a_comment_thread():
    s = new()
    s.run("review.newComment", cell="B2", text="Is this final?", author="Amy")
    s.run("review.replyComment", cell="B2", text="Yes.", author="Bo")
    s.run("review.resolveComment", cell="B2")
    c = s.doc().sheet.comments[(1, 1)]
    assert c.threaded and c.resolved and c.text == "Is this final?" and c.replies == (("Bo", "Yes."),)
    s.run("review.newNote", cell="C2", text="Checked")
    # Neither replaces the other.
    with pytest.raises(EngineError, match="already has a comment"):
        s.run("review.newNote", cell="B2", text="A note")
    with pytest.raises(EngineError, match="already has a note"):
        s.run("review.newComment", cell="C2", text="Why?")
    assert s.doc().sheet.comments[(1, 1)].replies == (("Bo", "Yes."),) and s.doc().sheet.comments[(1, 2)].text == "Checked"
    # Comments and replies have text.
    with pytest.raises(EngineError, match="Type a comment"):
        s.run("review.replyComment", cell="B2", text="  ")
    with pytest.raises(EngineError, match="Type a comment"):
        s.run("review.newComment", cell="D2", text="")
    s.run("review.deleteComment", cell="C2")
    s.run("review.newComment", cell="C2", text="Why?")
    assert s.doc().sheet.comments[(1, 2)].threaded


def test_viewport_sends_links_with_their_screentips():
    s = new()
    s.run("insert.hyperlink", cell="A1", target="https://example.com", text="site", tooltip="Our site")
    s.run("insert.hyperlink", cell="A2", target="Sheet1!K40")
    links = viewport(s.doc().wb, 0, RangeRef(0, 0, 30, 10))["links"]
    assert sorted((x["r"], x["c"], x["target"], x["tooltip"]) for x in links) == [(0, 0, "https://example.com", "Our site"), (1, 0, "Sheet1!K40", None)]


def test_dynamic_arrays_written_as_excel_does():
    s = new()
    s.run("cell.set", cell="A1", input="=SEQUENCE(2,2)")
    s.run("cell.set", cell="D1", input="=UNIQUE({1;1;2})")
    raw = base64.b64decode(s.run("file.saveBytes")["data"])
    import zipfile

    z = zipfile.ZipFile(io.BytesIO(raw))
    xml = z.read("xl/worksheets/sheet1.xml").decode()
    assert '<c r="A1" cm="1"><f t="array" ref="A1:B2">_xlfn.SEQUENCE(2,2)</f>' in xml
    assert "_xlfn.UNIQUE" in xml
    assert "xl/metadata.xml" in z.namelist()
    vals = openpyxl.load_workbook(io.BytesIO(raw), data_only=True).active
    assert [vals["A1"].value, vals["B2"].value] == [1, 4]


def test_csv_roundtrip():
    s = new()
    s.run("range.setValues", range="A1", values=[["a, b", 1.5, "=B1*2"], ["x\"y", True, None]])
    r = s.run("file.saveBytes", format="csv")
    text = base64.b64decode(r["data"]).decode("utf-8-sig")
    assert text.splitlines()[0] == '"a, b",1.5,3'
    s2 = Session()
    s2.run("file.open", data=r["data"], name="t.csv")
    assert val(s2, "A1") == "a, b" and val(s2, "C1") == 3 and val(s2, "A2") == 'x"y'


# ---------------------------------------------------------------- what-if, html


def test_goal_seek():
    s = new()
    s.run("range.setValues", range="A1", values=[[0.05], [10], [1000], ["=PMT(A1/12,A2*12,-A3)"]])
    r = s.run("data.goalSeek", setCell="A4", toValue=12, byChangingCell="A3")
    assert r["found"] and abs(val(s, "A4") - 12) < 0.001
    s.run("edit.undo")
    assert val(s, "A3") == 1000
    with pytest.raises(EngineError):
        s.run("data.goalSeek", setCell="A3", toValue=1, byChangingCell="A1")


def test_html_export_carries_formatting():
    s = new()
    s.run("range.setValues", range="A1", values=[["Title", None], [1234.5, "x"]])
    s.run("home.bold", range="A1")
    s.run("home.mergeCenter", range="A1:B1")
    s.run("home.numberFormat", range="A2", format="currency")
    s.run("pageLayout.orientation", value="landscape")
    page = s.run("file.exportHtml")["html"]
    assert "colspan='2'" in page and "font-weight:bold" in page
    assert "$1,234.50" in page and "size: letter landscape" in page


def _png(w: int = 2, h: int = 1) -> bytes:
    import struct
    import zlib

    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)

    rows = b"".join(b"\x00" + b"\xff\x00\x00" * w for _ in range(h))
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b"")


def test_picture_roundtrip_and_move():
    s = new()
    r = s.run("insert.picture", data=_png(4, 2), name="red.png", cell="C3")
    assert (r["width"], r["height"]) == (4, 2)
    s.run("object.move", id=r["id"], width=40, height=20, row=5, col=1)
    s2 = roundtrip(s)
    im = s2.doc().sheet.images[0]
    assert im.data == _png(4, 2) and im.mime == "image/png"
    assert (im.anchor.row, im.anchor.col) == (5, 1) and round(im.anchor.width) == 40
    assert objects(s2.doc().wb, 0)["images"][0]["src"].startswith("data:image/png;base64,")
    # The drawing holds a picture whose blip points at the media part, as Excel expects.
    import zipfile

    z = zipfile.ZipFile(io.BytesIO(base64.b64decode(s.run("file.saveBytes")["data"])))
    drawing = z.read("xl/drawings/drawing1.xml").decode()
    rels = z.read("xl/drawings/_rels/drawing1.xml.rels").decode()
    assert "<xdr:pic>" in drawing and 'r:embed="rId1"' in drawing
    assert 'Target="../media/image1.png"' in rels and z.read("xl/media/image1.png") == _png(4, 2)


# ---------------------------------------------------------------- spills and edits of many cells


def test_spill_anchor_follows_edits_of_many_cells():
    """A #SPILL! anchor is blocked, and freed again, by edits of many cells at once (a paste,
    Clear) as by one cell: the look for anchors near the changed cells is one pass per edit."""
    s = new()
    s.run("cell.set", cell="A1", input="=SEQUENCE(5)")
    assert [val(s, f"A{i}") for i in range(1, 6)] == [1, 2, 3, 4, 5]
    # Pasting onto the spill range blocks the anchor.
    s.run("edit.paste", text="x\ny\n", at="A3")
    assert val(s, "A1") == "#SPILL!" and val(s, "A3") == "x"
    # Pasting empty cells over the blockers frees it: one edit of several cells.
    s.run("edit.paste", text="\n\n", at="A3")
    assert [val(s, f"A{i}") for i in range(1, 6)] == [1, 2, 3, 4, 5]
    # Clear over a block that sits down and right of its anchor.
    s.run("cell.set", cell="B3", input="=SEQUENCE(1,4)")
    assert [val(s, c + "3") for c in "BCDE"] == [1, 2, 3, 4]
    s.run("edit.paste", text="1\t2\t3\t4\n5\t6\t7\t8\n", at="C2")
    assert val(s, "B3") == "#SPILL!"
    s.run("edit.clearContents", range="C2:F3")
    assert [val(s, c + "3") for c in "BCDE"] == [1, 2, 3, 4]


def test_spill_anchor_found_by_small_edits_on_a_big_sheet():
    """On a sheet with many cells, a change of one or two cells looks around those cells only
    (the other way of looking), and still finds the anchor it frees."""
    s = new()
    rows, cols = 100, 100
    s.run("edit.paste", text="\n".join("\t".join("7" for _ in range(cols)) for _ in range(rows)) + "\n", at="K1")
    s.run("cell.set", cell="A1", input="=SEQUENCE(5)")
    s.run("cell.set", cell="A3", input="x")
    assert val(s, "A1") == "#SPILL!"
    s.run("cell.set", cell="A3", input="")
    assert val(s, "A5") == 5
    s.run("edit.paste", text="x\ny\n", at="A3")
    assert val(s, "A1") == "#SPILL!"
    s.run("edit.paste", text="\n\n", at="A3")
    assert val(s, "A5") == 5


def test_pasting_many_cells_is_fast():
    """A paste costs about the same per cell as entering one cell: 50,000 cells took 7 seconds
    when every pasted cell scanned its own 65×65 window for #SPILL! anchors; well under a
    second now."""
    s = new()
    rows, cols = 500, 100
    text = "\n".join("\t".join(str(r * cols + c) for c in range(cols)) for r in range(rows)) + "\n"
    t0 = time.perf_counter()
    s.run("edit.paste", text=text)
    took = time.perf_counter() - t0
    assert val(s, "CV500") == rows * cols - 1
    assert took < 3.0, f"pasting {rows * cols} cells took {took:.1f}s"


# ---------------------------------------------------------------- trying a formula without storing it


def test_evaluate_check_tries_a_formula_as_typed_without_storing_it():
    """``formulas.evaluate`` with ``check``: what a formula would show in a cell, refusing what
    typing it would refuse (the AI's formula suggestions are checked this way before they show)."""
    s = new()
    s.run("range.setValues", range="A1", values=[["Item", "Qty", "Price"], ["Pen", 3, 1.5], ["Ink", 5, 2]])
    s.run("home.numberFormat", range="B4", code='"$"#,##0.00')
    s.run("home.bold", range="E2:E3")  # formatted, still empty
    undo = len(s.doc().undo)
    r = s.run("formulas.evaluate", formula="=SUMPRODUCT(B2:B3,C2:C3)", cell="B4", check=True)
    assert r["formula"] == "=SUMPRODUCT(B2:B3,C2:C3)" and r["value"] == 14.5 and r["shown"] == "$14.50"
    assert r["error"] is None and r["spill"] is None
    assert s.run("formulas.evaluate", formula="=TODAY()", cell="D9", check=True)["shown"].count("/") == 2
    assert s.run("formulas.evaluate", formula="=SEQUENCE(3)", cell="D1", check=True)["spill"] == [3, 1]
    assert s.run("formulas.evaluate", formula="=1/0", cell="D1", check=True)["error"] == "#DIV/0!"
    assert s.run("formulas.evaluate", formula="sum(b2:b3)", cell="D1", check=True)["formula"] == "=SUM(B2:B3)"
    assert s.run("formulas.evaluate", formula="=E2-E3", cell="B4", check=True)["blankRefs"] is True
    assert s.run("formulas.evaluate", formula="=B2-E3", cell="B4", check=True)["blankRefs"] is False
    assert s.run("formulas.evaluate", formula="=TODAY()", cell="B4", check=True)["blankRefs"] is False
    for bad in ("=SUM(B2:B3", "=IF(1)", "=B4*2", "=SUM(B:B)", "=Sheet1!B4"):
        with pytest.raises(EngineError):
            s.run("formulas.evaluate", formula=bad, cell="B4", check=True)
    assert s.run("cell.get", cell="B4")["input"] == "" and len(s.doc().undo) == undo
