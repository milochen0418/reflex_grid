"""Bugs that GridCraft (https://github.com/storytold/gridcraft), whose design gridcore follows, fixed
in its issues and pull requests, checked here so they stay fixed in gridcore. Each test names the
GridCraft issue (#…) or pull request (PR …). The expected answers were checked in Microsoft Excel 16
for Mac (en-US); where Excel answered otherwise than the issue said, the test says so and follows
Excel."""

import base64
import io
import math
import os
import re
import tempfile
import time
import warnings
import zipfile

import openpyxl
import pytest

from gridcore.calc import evaluate
from gridcore.core.addr import RangeRef
from gridcore.core.date import DateSystem
from gridcore.core.value import XlError, display
from gridcore.engine import EngineError, Session, viewport
from gridcore.model.workbook import Workbook
from gridcore.numfmt import text_function


def new():
    s = Session()
    s.new_workbook()
    return s


def val(s, a):
    return s.run("cell.get", cell=a)["value"]


def inp(s, a):
    return s.run("cell.get", cell=a)["input"]


def ev(formula, wb=None):
    wb = wb or Workbook()
    return display(evaluate(wb, wb.active_sheet, (0, 0), formula))


def cell(s, formula, at="Z1"):
    s.run("cell.set", cell=at, input=formula)
    return val(s, at)


def roundtrip(s, name="rt.xlsx"):
    s2 = Session()
    s2.run("file.open", data=s.run("file.saveBytes")["data"], name=name)
    return s2


def part(s, name):
    raw = base64.b64decode(s.run("file.saveBytes")["data"])
    return zipfile.ZipFile(io.BytesIO(raw)).read(name).decode()


def wb1904():
    wb = Workbook()
    wb.date_system = DateSystem.D1904
    return wb


def minimal_xlsx(sheet_body, workbook_extra=""):
    """A one-sheet package around ``sheet_body`` (the children of ``<worksheet>``)."""
    buf = io.BytesIO()
    z = zipfile.ZipFile(buf, "w")
    ct = "application/vnd.openxmlformats-officedocument.spreadsheetml"
    z.writestr(
        "[Content_Types].xml",
        '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        f'<Override PartName="/xl/workbook.xml" ContentType="{ct}.sheet.main+xml"/>'
        f'<Override PartName="/xl/worksheets/sheet1.xml" ContentType="{ct}.worksheet+xml"/>'
        f'<Override PartName="/xl/styles.xml" ContentType="{ct}.styles+xml"/></Types>',
    )
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    z.writestr(
        "_rels/.rels",
        f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rId1" Type="{rel}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
    )
    z.writestr(
        "xl/workbook.xml",
        f'<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="{rel}">'
        f'{workbook_extra}<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>',
    )
    z.writestr(
        "xl/_rels/workbook.xml.rels",
        f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        f'<Relationship Id="rId1" Type="{rel}/worksheet" Target="worksheets/sheet1.xml"/>'
        f'<Relationship Id="rId2" Type="{rel}/styles" Target="styles.xml"/></Relationships>',
    )
    z.writestr(
        "xl/styles.xml",
        '<?xml version="1.0" encoding="UTF-8"?><styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
        '<fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills>'
        '<borders count="1"><border/></borders><cellStyleXfs count="1"><xf/></cellStyleXfs><cellXfs count="1"><xf xfId="0"/></cellXfs>'
        '<dxfs count="1"><dxf><font><b/></font></dxf></dxfs></styleSheet>',
    )
    z.writestr(
        "xl/worksheets/sheet1.xml",
        f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="{rel}">{sheet_body}</worksheet>',
    )
    z.close()
    return base64.b64encode(buf.getvalue()).decode()


def open_data(data, name):
    s = Session()
    s.run("file.open", data=data, name=name)
    return s


# ---------------------------------------------------------------- functions


def test_search_positions_with_a_letter_whose_lower_case_is_longer():
    # #571: İ lower-cases to two characters; positions must still count it as one.
    assert ev('=SEARCH("b","İb")') == "2"
    assert ev('=SEARCH("B","İb")') == "2"
    assert ev('=MID("İb",SEARCH("b","İb"),1)') == "b"
    assert ev('=SEARCH("b?","aİbc")') == "3"
    assert ev('=SEARCH("b","aİbc",4)') == "#VALUE!"
    assert ev('=TEXTBEFORE("aİbXc","x",1,1)') == "aİb"


def test_case_insensitive_text_folds_one_letter_to_one_letter():
    # #571 had İ match i; Excel 16 for Mac folds by Unicode simple case folding: İ and ı are letters
    # of their own, ß is not ss; ς, σ and Σ are one letter, µ and μ too.
    for f, want in [('=SEARCH("i","aİbc")', "#VALUE!"), ('=SEARCH("I","aİb")', "#VALUE!"), ('=SEARCH("ı","aIb")', "#VALUE!"),
                    ('=SEARCH("ς","aΣb")', "2"), ('=SEARCH("µ","aΜb")', "2"), ('=SEARCH("ss","aßb")', "#VALUE!"),
                    ('=SEARCH("ß","aẞb")', "2"), ('="ς"="Σ"', "TRUE"), ('="İ"="i"', "FALSE"), ('=MATCH("i",{"İ","i"},0)', "2")]:
        assert ev(f) == want, f
    s = new()
    s.run("cell.set", cell="A1", input="İstanbul")
    assert cell(s, '=COUNTIF(A1,"i*")') == 0 and cell(s, '=COUNTIF(A1,"İSTANBUL")') == 1


def test_upper_and_lower_change_each_letter_into_one_letter():
    # Excel 16 for Mac: UPPER keeps ß and the letters whose upper case is two letters; LOWER("İ") is i.
    for f, want in [('=UPPER("straße")', "STRAßE"), ('=LOWER("İSTANBUL")', "istanbul"), ('=PROPER("ßa")', "ßa"), ('=UPPER("ŉ")', "ŉ"),
                    ('=LEN(UPPER("ﬀ"))', "1"), ('=UPPER("ǆ")', "Ǆ"), ('=PROPER("ǆa")', "Ǆa")]:
        assert ev(f) == want, f


def test_find_and_search_start_past_the_end():
    # #541 made both #VALUE!; Excel 16 for Mac lets FIND start just past the end of the text, not
    # SEARCH, which finds nothing in empty text.
    for f, want in [('=FIND("","abc",3)', "3"), ('=FIND("","abc",4)', "4"), ('=FIND("","abc",5)', "#VALUE!"), ('=FIND("","")', "1"),
                    ('=FIND("","a",2)', "2"), ('=FIND("x","abc",4)', "#VALUE!"), ('=FINDB("","abc",4)', "4"), ('=SEARCH("","abc",3)', "3"),
                    ('=SEARCH("","abc",4)', "#VALUE!"), ('=SEARCH("","")', "#VALUE!")]:
        assert ev(f) == want, f


def test_aggregate_needs_its_data_argument():
    # #572
    assert ev("=AGGREGATE(9,0)") == "#VALUE!"
    assert ev("=AGGREGATE(4,0)") == "#VALUE!"
    # Functions 1-13 take references only (Excel 16 for Mac); the array form (14-19) takes arrays.
    s = new()
    s.run("range.setValues", range="A1", values=[[1], [2]])
    for f, want in [("=AGGREGATE(9,0,{1,2})", "#VALUE!"), ("=AGGREGATE(9,0,A1:A2*1)", "#VALUE!"), ("=AGGREGATE(9,0,1)", "#VALUE!"),
                    ("=AGGREGATE(9,0,A1:A2)", 3), ("=AGGREGATE(9,0,OFFSET(A1,0,0,2,1))", 3), ("=AGGREGATE(14,6,{1,2,3},1)", 3),
                    ("=SUBTOTAL(9,{1,2})", "#VALUE!"), ("=SUBTOTAL(9,A1:A2)", 3)]:
        assert cell(s, f) == want, f


def test_subtotal_leaves_out_nested_subtotals_whatever_its_number():
    # #556
    s = new()
    s.run("cell.set", cell="A1", input="2")
    s.run("cell.set", cell="A2", input="=SUBTOTAL(9,A1)")
    assert cell(s, "=SUBTOTAL(109,A1:A2)") == 2
    assert cell(s, "=SUBTOTAL(9,A1:A2)") == 2
    assert cell(s, "=AGGREGATE(9,4,A1:A2)") == 4  # AGGREGATE keeps its own rule (options 4-7)


def test_prob_checks_only_that_the_probabilities_add_up_to_1():
    # #539 asked for #NUM! on a 0 probability (as Microsoft's documentation says); Excel 16 for Mac
    # takes any probabilities, -0.1 too, as long as they add up to 1 to 15 digits.
    for f, want in [("=PROB({1,2,3},{0,0.5,0.5},1,3)", "1"), ("=PROB({1,2},{-0.1,1.1},1)", "-0.1"), ("=PROB({1,2},{2,-1},2)", "-1"),
                    ("=PROB({1,2,3},{0.2,0.3,0.5},2,3)", "0.8"), ("=PROB({1,2},{0.5,0.6},1)", "#NUM!"),
                    ("=PROB({1,2,3},{0.2,0.30000000000001,0.5},1,3)", "#NUM!"), ("=PROB({1,2,3},{0.2,0.300000000000001,0.5},1,3)", "1")]:
        assert ev(f) == want, f


def test_rounding_at_precisions_past_the_number_range():
    # #540 and 4941bf8d: toward zero it is 0, away from zero 10^309 can't be stored.
    for f in ("ROUNDDOWN", "TRUNC", "ROUND"):
        for x in ("1234", "-1234", "0.5"):
            for d in ("-308", "-309", "-400"):
                assert ev(f"={f}({x},{d})") == "0", (f, x, d)
    assert ev("=ROUNDUP(1234,-309)") == "#NUM!"
    assert ev("=ROUNDUP(0,-309)") == "0"
    assert ev("=ROUNDUP(5E+307,-308)") == "1E+308"


def test_roman_forms():
    # #733, and Microsoft's ROMAN examples for 499.
    expected = {
        (499, 0): "CDXCIX", (499, 1): "LDVLIV", (499, 2): "XDIX", (499, 3): "VDIV", (499, 4): "ID",
        (99, 2): "IC", (199, 2): "CIC", (990, 2): "XM", (1999, 2): "MXMIX", (995, 3): "VM",
        (495, 2): "XDV", (495, 3): "VD", (995, 2): "XMV", (3999, 0): "MMMCMXCIX", (0, 0): "",
    }
    for (n, form), want in expected.items():
        assert ev(f"=ROMAN({n},{form})") == want, (n, form)
    for f, want in [("=ROMAN(4)", "IV"), ("=ROMAN(1999)", "MCMXCIX"), ("=ROMAN(499,)", "CDXCIX"), ('=ROMAN(499,"2")', "XDIX"), ("=ROMAN(A1)", "")]:
        assert ev(f) == want, f
    assert ev("=ROMAN(499,TRUE)") == "CDXCIX"
    assert ev("=ROMAN(499,FALSE)") == "ID"
    assert ev("=ROMAN(499,5)") == "#VALUE!"


def test_edate_and_eomonth_in_early_1900():
    # #734: serial 0 is day 0 of January and February 1900 has 28 days there.
    for f, want in [("=EDATE(0,0)", "0"), ("=EDATE(0,1)", "31"), ("=EDATE(0,12)", "366"), ("=EOMONTH(0,1)", "59"),
                    ("=EDATE(31,1)", "59"), ("=EOMONTH(59,0)", "59"), ("=EOMONTH(DATE(1900,1,15),1)", "59"),
                    ("=EDATE(60,1)", "89"), ("=EOMONTH(0,0)", "31"), ("=DATE(1900,2,29)", "60")]:
        assert ev(f) == want, f


def test_mode_skips_typed_text_and_logical_values():
    # #736
    for f, want in [('=MODE(1,1,"x")', "1"), ('=MODE(1,2,"x")', "#N/A"), ('=MODE(1,"1",2,3)', "#N/A"),
                    ('=MODE("1",1,#DIV/0!,2,3)', "#DIV/0!"), ("=MODE(1,TRUE,TRUE,2)", "#N/A"), ("=MODE(1,1,TRUE)", "1"),
                    ('=MODE.SNGL(3,3,"a",TRUE)', "3")]:
        assert ev(f) == want, f


def test_mirr_countif_indirect_and_usdollar():
    # #738
    assert ev("=MIRR({-10000,-3400,-6500,-1000},0.07,0.01)") == "-1"
    assert ev("=MIRR({100,200,300},0.07,0.01)") == "#DIV/0!"
    s = new()
    s.run("cell.set", cell="H3", input='=""')
    assert cell(s, '=COUNTIF(H3,"<=")') == 0
    assert cell(s, '=COUNTIF(H3,">=")') == 0
    s.run("cell.set", cell="I3", input="30")
    assert cell(s, '=INDIRECT("I3",FALSE)') == "#REF!"
    assert cell(s, '=INDIRECT("R3C9",FALSE)') == 30
    assert ev("=USDOLLAR(1)") == "$1.00"


def test_median_of_huge_numbers_doesnt_overflow():
    # #1579
    assert ev("=MEDIAN(1E+308,1E+308)") == "1E+308"
    assert ev("=MEDIAN(-1E+308,-1E+308)") == "-1E+308"
    assert ev("=MEDIAN(9E+307,9E+307)") == "9E+307"  # Excel takes no number past 9.99999999999999E+307 typed in
    s = new()
    s.run("range.setValues", range="A1", values=[[9e307], [9e307]])
    assert cell(s, "=AGGREGATE(12,6,A1:A2)") == 9e307
    assert ev("=MEDIAN(1,2,3,4)") == "2.5"


def test_criteria_longer_than_255_characters():
    # PR 686: #VALUE!. With several criteria pairs too, in any position (Excel 16 for Mac; PR 686
    # matched the first 255 units there).
    s = new()
    s.run("cell.set", cell="A1", input="a" * 300)
    s.run("cell.set", cell="B1", input="a" * 300)
    s.run("cell.set", cell="A2", input="a" * 255)
    for f in ("=COUNTIF(A1,B1)", "=SUMIF(A1,B1,A1)", "=AVERAGEIF(A1,B1,A1)", "=COUNTIFS(A1,B1)", "=COUNTIFS(A1:A2,B1,A1:A2,B1)",
              '=COUNTIFS(A1:A2,"<>",A1:A2,B1)', '=SUMIFS(A1:A2,A1:A2,"<>",A1:A2,B1)'):
        assert cell(s, f) == "#VALUE!", f
    assert cell(s, "=COUNTIF(A1,LEFT(B1,255))") == 0
    assert cell(s, "=COUNTIF(A2,LEFT(B1,255))") == 1
    assert cell(s, "=COUNTIFS(A1:A2,LEFT(B1,255),A1:A2,LEFT(B1,255))") == 1


def test_ifs_functions_with_several_criteria_pairs():
    # Found while checking PR 686: past the first pairs, a criteria range was taken for a single
    # value, so COUNTIFS with two pairs (SUMIFS with three) gave #VALUE!.
    s = new()
    s.run("range.setValues", range="A1", values=[[1, 1, 1], [2, 0, 1], [3, 5, 0]])
    for f, want in [('=COUNTIFS(A1:A3,">0",B1:B3,">0")', 2), ('=COUNTIFS(A1:A3,">0",B1:B3,">=0",C1:C3,1)', 2),
                    ('=SUMIFS(A1:A3,A1:A3,">0",B1:B3,">0")', 4), ('=SUMIFS(A1:A3,A1:A3,">0",B1:B3,">0",C1:C3,">0")', 1),
                    ('=AVERAGEIFS(A1:A3,A1:A3,">0",B1:B3,">0")', 2), ('=MAXIFS(A1:A3,B1:B3,">0",C1:C3,">0")', 1),
                    ('=MINIFS(A1:A3,B1:B3,">=0",C1:C3,">0")', 1), ('=SUM(COUNTIFS(A1:A3,{1,2,3}))', 3)]:
        assert cell(s, f) == want, f


def test_textbefore_and_textafter_search_from_the_end_for_negative_instances():
    # PR 1941
    for f, want in [('=TEXTBEFORE("aaa","aa",-1)', "a"), ('=TEXTAFTER("aaa","aa",-1)', ""), ('=TEXTBEFORE("a-b-c","-",-1)', "a-b"),
                    ('=TEXTAFTER("a-b-c","-",-2)', "b-c"), ('=TEXTBEFORE("xAyaz","A",-1,1)', "xAy"), ('=TEXTBEFORE("aaa","aa",1)', "")]:
        assert ev(f) == want, f


def test_month_name_dates_as_excel_reads_them():
    # #566 took "Mar 4 26" for a date; Excel 16 for Mac (en-US) doesn't: "month day year" needs a
    # comma after the day (then a space), "day month year" takes none. Without a year the day is in
    # this year, and a number that can't be a day of the month is the year (Mar 32 is March 1932).
    from gridcore.core.date import serial_from_ymd

    s = new()

    def typed(t):
        s.run("cell.set", cell="A1", input=t)
        return val(s, "A1")

    for t, want in [("Mar 4, 26", 46085), ("March 4, 26", 46085), ("Mar, 4, 2026", 46085), ("Mar 4 , 26", 46085), ("March 4,  2026", 46085),
                    ("4 March 26", 46085), ("4 Mar 2026", 46085), ("4-Mar-26", 46085), ("4/Mar/26", 46085), ("Mar 4, 99", 36223),
                    ("Mar 4, 5", 38415), ("Mar 32", 11749), ("Mar 99", 36220), ("Mar 0", 36586), ("Mar-32", 11749), ("Mar 1900", 61),
                    ("Jan 2024", 45292), ("March-2026", 46082)]:
        assert typed(t) == want, t
    for t in ("Mar 4 26", "March 4 26", "Mar 4 2026", "Mar 26 2026", "Sep 30 99", "Mar 4,2026", "4 Mar, 26", "4 March, 2026", "Mar. 4", "Mar.26",
              "Mar/4/26", "32 Mar", "0 Mar", "Mar 100", "Mar 1899", "Mar 4, 1776", "Feb 30, 26", "Tue, Mar 4, 2026", "March 4th, 2026"):
        assert typed(t) == t, t
    y = time.localtime().tm_year
    for t, (m, d) in [("Mar 4", (3, 4)), ("4 Mar", (3, 4)), ("mar 4", (3, 4)), ("Sept 4", (9, 4)), ("Mar 26", (3, 26)), ("Mar-26", (3, 26)),
                      ("Mar/4", (3, 4)), ("Jan 24", (1, 24))]:
        assert typed(t) == serial_from_ymd(DateSystem.D1900, y, m, d), t
    s.run("cell.set", cell="B1", input="Jan 24")
    assert s.run("cell.get", cell="B1")["text"] == "24-Jan"
    assert typed("Mar 4 1:30 PM") == serial_from_ymd(DateSystem.D1900, y, 3, 4) + 0.5625


def test_lambda_takes_one_argument_position_per_parameter():
    # #557
    for f, want in [("=LAMBDA(a,b,a+b)(2)", "#VALUE!"), ("=LAMBDA(a,b,a+b)(2,3)", "5"), ("=LAMBDA(a,b,a+b)(2,3,4)", "#VALUE!"),
                    ('=LAMBDA(a,b,IF(ISOMITTED(b),"Missing",a+b))(1,)', "Missing"), ("=LAMBDA(7)()", "7"), ("=LAMBDA(x,ISOMITTED(x))()", "TRUE")]:
        assert ev(f) == want, f
    s = new()
    s.run("formulas.defineName", name="AddPair", refersTo="=LAMBDA(a,b,a+b)")
    assert cell(s, "=AddPair(2)") == "#VALUE!"
    assert cell(s, "=AddPair(2,3)") == 5


def test_functions_passed_by_name():
    # #902 and PR 486: GROUPBY(…, SUM) as Microsoft documents it, BYROW(…, SUM), MAP(…, ABS).
    s = new()
    s.run("range.setValues", range="A1", values=[["K", "V"], ["a", 1], ["b", 2], ["a", 3]])
    s.run("cell.set", cell="D1", input="=GROUPBY(A2:A4,B2:B4,SUM)")
    assert [val(s, a) for a in ("D1", "E1", "D2", "E2")] == ["a", 4, "b", 2]
    assert cell(s, "=BYROW(B2:B4,SUM)", "G1") == 1 and val(s, "G3") == 3
    assert cell(s, "=MAP(-B2:B4,ABS)", "H1") == 1
    assert cell(s, "=REDUCE(0,B2:B4,SUM)", "I1") == 6
    assert cell(s, "=SCAN(0,B2:B4,SUM)", "J1") == 1 and val(s, "J3") == 6
    assert cell(s, "=MAKEARRAY(2,2,PRODUCT)", "K1") == 1 and val(s, "L2") == 4
    # A LET variable or a defined name keeps its meaning.
    assert cell(s, "=LET(SUM,LAMBDA(x,x*10),BYROW(B2:B3,SUM))", "M1") == 10
    # Files spell such names with _xleta. (Excel's form) and read back unchanged.
    x = part(s, "xl/worksheets/sheet1.xml")
    assert "_xlfn.GROUPBY(A2:A4,B2:B4,_xleta.SUM)" in x and "_xlfn.BYROW(B2:B4,_xleta.SUM)" in x
    s2 = roundtrip(s)
    assert inp(s2, "D1") == "=GROUPBY(A2:A4,B2:B4,SUM)" and val(s2, "E1") == 4


def test_number_formats_drop_the_sign_of_a_negative_that_rounds_to_zero():
    # #729, Excel 365's table.
    for code, v, want in [("0.00", -0.0033, "0.00"), ("0.00", -2.89e-09, "0.00"), ("0.00", -0.005, "-0.01"), ("0", -0.4, "0"),
                          ("0", -0.5, "-1"), ("0.0", -0.004, "0.0"), ("#,##0.00", -0.004, "0.00"), ("0.00%", -2.89e-09, "0.00%"),
                          ("0.00%", -0.0033, "-0.33%"), ("0.00;(0.00)", -0.0033, "(0.00)"), ('0.00;-0.00;"zero"', -0.0033, "-0.00"),
                          ("0.00E+00", -2.89e-09, "-2.89E-09")]:
        assert text_function(v, code) == want, (code, v)


def test_scientific_formats():
    # #768, in the forms Excel 16 for Mac takes (its answers): % and commas don't scale an exponent,
    # zero fills # placeholders. A lowercase e is the era year to Excel, so those codes are refused.
    for code, v, want in [("#E+#%", 12345, "1E+4%"), ("#,E+#", 12345, "1E+4"), ("#%E-#", 12345, "1%E4"), ("####.####E+#", 12345, "1.2345E+4"),
                          ("#.#E+#", -1.23456789e-05, "-1.2E-5"), ("#.#E+#", 0, "0.E+0"), ("####.####E+#", 0, "0000.E+0"),
                          ("0.00E+00", 12345, "1.23E+04"), ("0.00E+00", 0, "0.00E+00"), ("##0.0E+0", 12345, "12.3E+3"),
                          ("0.0E+0%", 12345, "1.2E+4%"), ("0.0E+0%", 0.004, "4.0E-3%"), ("0.0E-0", 12345, "1.2E4")]:
        assert text_function(v, code) == want, (code, v)
    for code in ("0.00e+00", "#,e+#", "####.####e+#", "0E+", "E+0", "0.0E0", "0.0E+0E+0"):
        with pytest.raises(XlError):
            text_function(12345, code)


def test_format_codes_excel_refuses():
    # #807 (past 255 characters) and #810 (an open [), and the rest of what Excel 16 for Mac refuses
    # as a number format: TEXT() is #VALUE! and Format Cells says so.
    bad = ["0" * 256, "[0;0", "[0", "0[", '0"', "0\\", "0*", "0_", "0;0;0;@;0", "[xyz]0", "[Purple]0", "[Color57]0", "[ENG]0", "[d]", "[h]:[m]",
           "[Red][Blue]0", "[]0", "0 yyyy", "d 0", "yyyy%", "h:mm:ss.0#", "# ?/? d", "@ h", "0 @", "General 0"]
    good = ["0" * 255, "0 x", "[Red]0", "[color 10]0", "[DBNum1]0", "[$€-x-euro2]0", "[>=1][Red]0", "mm:ss.0", "hh.0", "[s].00", "d.m.yyyy",
            "h:mm,", "0 AM/PM", "General h", "0;0;0;@", ";;;", "@@", "0 _)", '_(* #,##0.00_);_(* (#,##0.00);_(* "-"??_);_(@_)']
    for code in bad:
        with pytest.raises(XlError):
            text_function(1, code)
    for code in good:
        text_function(1, code)
    s = new()
    s.run("cell.set", cell="A1", input="5")
    with pytest.raises(EngineError):
        s.run("home.numberFormat", range="A1", code="[0;0")
    with pytest.raises(EngineError):
        s.run("home.formatCells", range="A1", numFmt="0" * 256)
    s.run("home.numberFormat", range="A1", code="0.00")
    assert s.run("cell.get", cell="A1")["text"] == "5.00"


# ---------------------------------------------------------------- calculation


def test_cells_that_only_hold_a_format_dont_block_a_spill():
    # #1011 (and the bold empty cell before it): only content blocks a spill, as in Excel.
    s = new()
    s.run("home.bold", range="C1")
    s.run("cell.set", cell="A1", input="=SEQUENCE(1,3)")
    assert val(s, "A1") == 1 and val(s, "C1") == 3
    # A styled blocker far from the anchor, cleared: the spill comes back.
    s.run("cell.set", cell="A3", input="=SEQUENCE(1,80)")
    s.run("cell.set", cell="CB3", input="x")
    s.run("home.bold", range="CB3")
    assert val(s, "A3") == "#SPILL!"
    s.run("edit.clearContents", range="CB3")
    assert val(s, "A3") == 1 and val(s, "CB3") == 80


def test_a_formula_reading_a_cleared_spill_past_65536_cells_recalculates():
    # #558
    s = new()
    s.run("cell.set", cell="A1", input="=SEQUENCE(65537)")
    s.run("cell.set", cell="C2", input="=A65537")
    assert val(s, "C2") == 65537
    s.run("cell.set", cell="A1", input="5")
    assert val(s, "C2") == 0


def test_subtotal_follows_rows_hidden_shown_and_filtered():
    # PR 796, and filters: SUBTOTAL and AGGREGATE look again whenever rows are hidden or shown.
    s = new()
    s.run("range.setValues", range="A1", values=[["H"], [1], [2], [3]])
    s.run("cell.set", cell="C1", input="=SUBTOTAL(109,A2:A4)")
    s.run("cell.set", cell="C2", input="=SUBTOTAL(9,A2:A4)")
    s.run("cell.set", cell="C3", input="=AGGREGATE(9,5,A2:A4)")
    s.run("home.hideRows", rows="3:3")
    assert [val(s, a) for a in ("C1", "C2", "C3")] == [4, 6, 4]
    s.run("home.unhideRows", rows="2:4")
    assert val(s, "C1") == 6
    s.run("home.rowHeight", rows="3:3", height=0)
    assert val(s, "C1") == 4
    s.run("home.rowHeight", rows="3:3", height=20)
    assert val(s, "C1") == 6
    s.run("data.filter", range="A1:A4")
    s.run("data.filterColumn", col="A", criterion={"kind": "values", "values": ["1"]})
    assert [val(s, a) for a in ("C1", "C2", "C3")] == [1, 1, 1]
    s.run("data.clearFilter")
    assert [val(s, a) for a in ("C1", "C2", "C3")] == [6, 6, 6]


def test_a_table_total_row_follows_its_filter():
    s = new()
    s.run("range.setValues", range="A1", values=[["Item", "Qty"], ["a", 2], ["b", 3], ["c", 5]])
    s.run("insert.table", range="A1:B4", name="T")
    s.run("table.totalRow", on=True, table="T")
    assert inp(s, "B5") == "=SUBTOTAL(109,T[Qty])" and val(s, "B5") == 10
    s.run("selection.set", range="B2")
    s.run("data.filterColumn", col="A", criterion={"kind": "values", "values": ["a", "c"]})
    assert val(s, "B5") == 7


def test_long_operator_chains_calculate():
    # PR 484 and #527: Excel allows them up to its 8,192 characters.
    s = new()
    s.run("cell.set", cell="A1", input="1")
    chain = "=" + "+".join(["A1"] * 2000)
    assert cell(s, chain, "C1") == 2000
    assert cell(s, "=" + "&".join(['"a"'] * 1500), "C2") == "a" * 1500
    assert cell(s, "=" + "*".join(["1"] * 3000) + "=1", "C3") is True
    s.run("home.insertRows", rows="1:1")
    assert inp(s, "C2") == "=" + "+".join(["A2"] * 2000) and val(s, "C2") == 2000
    assert val(roundtrip(s), "C2") == 2000


def test_formula_text_is_at_most_8192_characters():
    # PR 1330: longer text, postfix % nesting past the limit: refused, never a crash.
    s = new()
    ok = "=" + "+".join(["1"] * 4096)  # 8,192 characters with the = (Excel 16 for Mac counts it)
    assert cell(s, ok, "A1") == 4096
    for bad in ("=" + "+".join(["1"] * 4095) + "+12", "=" + "+".join(["1"] * 4097), "=1" + "%" * 5000, "=" + "(" * 300 + "1" + ")" * 300):
        with pytest.raises(EngineError):
            s.run("cell.set", cell="A2", input=bad)
    assert cell(s, "=1" + "%" * 3, "A3") == pytest.approx(1e-6)


def test_today_and_now_in_a_1904_workbook():
    # #578
    wb = wb1904()
    assert ev("=YEAR(TODAY())", wb) == str(time.localtime().tm_year)
    assert ev("=TODAY()-DATE(YEAR(TODAY()),MONTH(TODAY()),DAY(TODAY()))", wb) == "0"
    assert ev("=INT(NOW())=TODAY()", wb) == "TRUE"


def test_date_text_in_a_1904_workbook():
    # #564 and #565
    wb = wb1904()
    assert ev('="1/1/1904"+0', wb) == "0"
    assert ev('=VALUE("1/2/1904")', wb) == "1"
    assert ev('=-"1/2/1904"', wb) == "-1"
    assert ev('="1/2/1904"+0') == "1463"  # 1900 workbooks unchanged
    assert ev('=TEXT(2957003,"m/d/yyyy")', wb) == "12/31/9999"
    assert ev('=TEXT(2957004,"m/d/yyyy")', wb) == "#VALUE!"
    assert ev('=TEXT(2958465,"m/d/yyyy")') == "12/31/9999"


def test_negative_dates_and_times_in_a_1904_workbook():
    # Excel 16 for Mac shows them with a minus sign in a 1904 workbook; a 1900 one can't (#VALUE!).
    wb = wb1904()
    for f, want in [('=TEXT(-1,"m/d/yyyy")', "-1/2/1904"), ('=TEXT(-1.5,"m/d/yyyy h:mm")', "-1/2/1904 12:00"), ('=TEXT(-0.25,"h:mm")', "-6:00"),
                    ('=TEXT(-1.5,"[h]:mm")', "-36:00")]:
        assert ev(f, wb) == want, f
    assert ev('=TEXT(-1,"m/d/yyyy")') == "#VALUE!"


def test_index_area_number_below_1():
    # PR 476 (an area number past the last area is #REF!), and one below 1 is #VALUE! (Excel 16 for
    # Mac; 0.5 counts as 0).
    s = new()
    s.run("range.setValues", range="A1", values=[[1, 2, 5, 6], [3, 4, 7, 8]])
    for f, want in [("=INDEX((A1:B2,C1:D2),1,1,2)", 5), ("=INDEX((A1:B2,C1:D2),1,1,3)", "#REF!"), ("=INDEX(A1:B2,1,1,1E+300)", "#REF!"),
                    ("=INDEX(A1:B2,1,1,-1E+300)", "#VALUE!"), ("=INDEX(A1:B2,1,1,-1)", "#VALUE!"), ("=INDEX(A1:B2,1,1,0)", "#VALUE!"),
                    ("=INDEX(A1:B2,1,1,0.5)", "#VALUE!"), ("=INDEX(A1:B2,1,1,1.9)", 1)]:
        assert cell(s, f) == want, f


def test_whole_rows_and_columns_take_one_dollar_sign():
    # PR 1922: $$1:$3 is no reference (Excel 16 for Mac: #REF! from INDIRECT).
    s = new()
    for f, want in [('=ROWS(INDIRECT("$$1:$3"))', "#REF!"), ('=ROWS(INDIRECT("$1:$$3"))', "#REF!"), ('=COLUMNS(INDIRECT("$$A:$B"))', "#REF!"),
                    ('=ROWS(INDIRECT("+1:+3"))', "#REF!"), ('=ROWS(INDIRECT("$1:$3"))', 3), ('=COLUMNS(INDIRECT("$A:B"))', 2)]:
        assert cell(s, f) == want, f
    assert RangeRef.parse("$$1:$3") is None and RangeRef.parse("$1:$3") is not None


# ---------------------------------------------------------------- formula text and files


def test_an_intersection_with_a_deleted_operand_reads_back():
    # #568
    s = new()
    for f in ("=A1:B2 #REF!", "=#REF! A1:B2", "=(A1:B2) #REF!"):
        assert cell(s, f, "H1") == "#REF!" and inp(s, "H1") == f
    s.run("cell.set", cell="H2", input="=SUM(A1:B2 D1:E2)")
    s.run("home.deleteColumns", cols="D:E")
    assert inp(s, "F2") == "=SUM(A1:B2 #REF!)"
    s2 = roundtrip(s)
    s2.run("home.insertColumns", cols="A:A")
    assert inp(s2, "G2") == "=SUM(B1:C2 #REF!)"


def test_sheet_scoped_names_keep_their_sheet_quoted():
    # #567, and renames to names that need quotes.
    s = new()
    s.run("home.insertSheet", name="My Sheet", index=1)
    s.run("sheet.activate", sheet="My Sheet")
    s.run("cell.set", cell="A1", input="42")
    s.run("formulas.defineName", name="Total", refersTo="='My Sheet'!$A$1", scope="My Sheet")
    s.run("sheet.activate", sheet="Sheet1")
    assert cell(s, "='My Sheet'!Total", "B1") == 42 and inp(s, "B1") == "='My Sheet'!Total"
    s2 = roundtrip(s)
    s2.run("sheet.activate", sheet="Sheet1")
    assert val(s2, "B1") == 42
    for name, shown in (("It's", "='It''s'!Total"), ("Plain", "=Plain!Total")):
        cur = s.doc().wb.sheets[1].name
        s.run("sheet.rename", sheet=cur, name=name)
        s.run("sheet.activate", sheet="Sheet1")
        assert inp(s, "B1") == shown and val(s, "B1") == 42


def test_user_defined_prefix_stays():
    # #726: _xludf. marks a user-defined function, not the built-in of the same name.
    s = new()
    s.run("cell.set", cell="A1", input='=_xludf.CONCAT("a","b")')
    assert val(s, "A1") == "#NAME?" and inp(s, "A1") == '=_xludf.CONCAT("a","b")'
    assert '_xludf.CONCAT("a","b")' in part(s, "xl/worksheets/sheet1.xml")
    assert cell(s, '=CONCAT("a","b")', "A2") == "ab"


def test_table_columns_keep_spaces_at_their_edges():
    # #732, Excel 365's table: the column is "Qty ", and Table1[Qty] names no column.
    s = new()
    s.run("range.setValues", range="A1", values=[["Item", "Qty "], ["a", 2], ["b", 3]])
    s.run("insert.table", range="A1:B3", name="Table1")
    s.run("table.totalRow", on=True, table="Table1")
    assert val(s, "B1") == "Qty "
    assert cell(s, "=SUM(Table1[Qty ])", "F1") == 5 and inp(s, "F1") == "=SUM(Table1[[Qty ]])"
    assert cell(s, "=SUM(Table1[[Qty ]])", "F2") == 5
    assert cell(s, "=Table1[[#Totals],[Qty ]]", "F3") == 5
    assert cell(s, "=SUBTOTAL(109,Table1[Qty ])", "F4") == 5
    assert cell(s, "=SUM(Table1[Qty])", "F5") == "#REF!"
    assert inp(s, "B4") == "=SUBTOTAL(109,Table1[[Qty ]])"
    t = part(s, "xl/tables/table1.xml")
    assert 'name="Qty "' in t
    s2 = roundtrip(s)
    assert val(s2, "F1") == 5 and val(s2, "B4") == 5


def test_a_workbook_opens_by_its_content_not_its_name():
    # #569
    s = new()
    s.run("cell.set", cell="A1", input="5")
    data = s.run("file.saveBytes")["data"]
    for name in ("book.csv", "book.ods", "book.dat"):
        assert val(open_data(data, name), "A1") == 5, name
    assert val(open_data(base64.b64encode(b"a,b\r\n1,2\r\n").decode(), "data.xlsx"), "B2") == 2


def test_document_properties_keep_their_spaces():
    # #554
    s = new()
    props = {"title": "  Title  ", "subject": "\n", "author": "\tA ", "company": " C ", "keywords": " k ", "description": "\n d \n"}
    s.run("file.properties", **props)
    p = roundtrip(s).run("file.properties")
    assert {k: p.get(k) for k in props} == props


def test_a_table_always_has_a_data_row():
    # #511 saved a table of a header and a total row only; Excel 16 for Mac finds that file damaged,
    # and one with a header row only too. As in Excel: a table made from a header row takes the row
    # below as its data row, and deleting all its data rows keeps the first, emptied.
    s = new()
    s.run("range.setValues", range="A1", values=[["Name", "Qty"]])
    s.run("insert.table", range="A1:B1", name="Tab")
    s.run("table.totalRow", on=True, table="Tab")
    t = roundtrip(s).doc().sheet.tables
    assert [(x.name, x.range.a1(), x.totals_row) for x in t] == [("Tab", "A1:B3", True)]
    with pytest.raises(EngineError):
        s.run("table.resize", range="A1:B2", table="Tab")
    s = new()
    s.run("range.setValues", range="A1", values=[["Name", "Qty"], ["a", 1], ["b", 2], ["c", 3]])
    s.run("range.setValues", range="D1", values=[["d1"], ["d2"], ["d3"], ["d4"], ["d5"]])
    s.run("insert.table", range="A1:B4", name="Tab")
    s.run("table.totalRow", on=True, table="Tab")
    s.run("home.deleteRows", rows="2:4")
    t = s.doc().sheet.tables[0]
    assert t.range.a1() == "A1:B3" and t.totals_row
    assert [val(s, a) for a in ("A2", "B2", "D2", "D3")] == [None, None, "d2", "d5"]
    assert inp(s, "B3") == "=SUBTOTAL(109,Tab[Qty])"
    # A table without a data row (from another program) isn't written.
    from gridcore.model.features import Table, TableColumn
    from gridcore.xlsx import write_xlsx

    wb = Workbook()
    wb.sheets[0].tables.append(Table(1, "Bare", RangeRef(0, 0, 1, 1), totals_row=True, columns=[TableColumn("A"), TableColumn("B")]))
    assert not [n for n in zipfile.ZipFile(io.BytesIO(write_xlsx(wb))).namelist() if n.startswith("xl/tables/")]


def test_every_range_of_a_conditional_format_and_validation_is_read():
    # #512
    sq = " ".join(f"A{2 * i + 1}" for i in range(12000))
    body = (
        "<sheetData/>"
        f'<conditionalFormatting sqref="{sq}"><cfRule type="expression" dxfId="0" priority="1"><formula>TRUE</formula></cfRule></conditionalFormatting>'
        f'<dataValidations count="1"><dataValidation type="whole" sqref="{sq}"><formula1>1</formula1></dataValidation></dataValidations>'
    )
    sh = open_data(minimal_xlsx(body), "x.xlsx").doc().sheet
    assert len(sh.cond_formats[0].ranges) == 12000 and len(sh.validations[0].ranges) == 12000


def test_hyperlinks_past_the_cell_cap_are_reported():
    # #513
    s = open_data(minimal_xlsx('<sheetData/><hyperlinks><hyperlink ref="A1:A12000" location="Sheet1!B1"/></hyperlinks>'), "x.xlsx")
    assert len(s.doc().sheet.hyperlinks) == 10000
    assert any("hyperlinks past the first 10,000 cells were dropped" in w for w in s.doc().warnings)
    s = open_data(minimal_xlsx('<sheetData/><hyperlinks><hyperlink ref="A1:A3" location="Sheet1!B1"/></hyperlinks>'), "x.xlsx")
    assert len(s.doc().sheet.hyperlinks) == 3 and not s.doc().warnings


def test_a_link_to_another_file_stays_a_link_to_that_file():
    # #800: files are external relationships, places in the workbook are locations.
    s = new()
    s.run("home.insertSheet", name="Sheet2", index=1)
    s.run("sheet.activate", sheet="Sheet1")
    for c, target in (("A1", "Source.xlsx"), ("A2", "..\\Reports\\Q1.xlsx"), ("A3", "https://example.com"), ("A4", "Sheet2!B5"), ("A5", "B5")):
        s.run("insert.hyperlink", cell=c, target=target)
    x = part(s, "xl/worksheets/sheet1.xml")
    rels = part(s, "xl/worksheets/_rels/sheet1.xml.rels")
    assert 'location="Source.xlsx"' not in x and 'Target="Source.xlsx"' in rels and "Q1.xlsx" in rels
    assert 'location="Sheet2!B5"' in x and 'location="B5"' in x
    links = roundtrip(s).doc().sheet.hyperlinks
    assert links[(0, 0)].external and links[(0, 0)].target == "Source.xlsx"
    assert not links[(3, 0)].external


def test_links_open_only_web_and_mail_addresses():
    # PR 780 (#777): a workbook's link can't run script in the browser or reach other schemes.
    s = new()
    s.run("home.insertSheet", name="Sheet2", index=1)
    s.run("sheet.activate", sheet="Sheet1")
    for target in ("javascript:alert(1)//x.y", " javascript:alert(1)", "JAVASCRIPT:x", "data:text/html,x", "file:///etc/passwd",
                   "vbscript:x", "java\tscript:alert(1)", "Source.xlsx"):
        s.run("insert.hyperlink", cell="A1", target=target)
        with pytest.raises(EngineError):
            s.run("hyperlink.follow", cell="A1")
    for target, url in (("https://example.com/a", "https://example.com/a"), ("mailto:a@b.co", "mailto:a@b.co"),
                        ("www.example.com", "http://www.example.com")):
        s.run("insert.hyperlink", cell="A1", target=target)
        assert s.run("hyperlink.follow", cell="A1") == {"url": url}
    s.run("insert.hyperlink", cell="A1", target="Sheet2!B5")
    assert s.run("hyperlink.follow", cell="A1") == {"sheet": "Sheet2", "range": "B5"}
    # A range on the same sheet is a place, not a link with the scheme `B5`.
    s.run("sheet.activate", sheet="Sheet1")
    for target, rng in (("B5:C9", "B5:C9"), ("#D3:E4", "D3:E4"), ("AB10:AC12", "AB10:AC12")):
        s.run("insert.hyperlink", cell="A1", target=target)
        assert s.run("hyperlink.follow", cell="A1") == {"sheet": "Sheet1", "range": rng}
    # A link read from a file goes through the same check.
    body = '<sheetData/><hyperlinks><hyperlink ref="A1" r:id="rId9"/></hyperlinks>'
    data = base64.b64decode(minimal_xlsx(body))
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(buf, "w") as zout:
        for n in zin.namelist():
            zout.writestr(n, zin.read(n))
        zout.writestr(
            "xl/worksheets/_rels/sheet1.xml.rels",
            '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId9" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
            'Target="javascript:alert(document.cookie)//x.y" TargetMode="External"/></Relationships>',
        )
    s2 = open_data(base64.b64encode(buf.getvalue()).decode(), "x.xlsx")
    with pytest.raises(EngineError):
        s2.run("hyperlink.follow", cell="A1")


def test_a_workbook_with_a_password_to_open_is_not_called_an_xls_file():
    # #11: such a workbook is a compound file holding the encrypted package.
    from gridcore.xlsx.cfb import Storage, write_cfb

    root = Storage()
    root.set_stream("EncryptionInfo", b"\x04\x00\x04\x00" + b"x" * 100)
    root.set_stream("EncryptedPackage", b"\x00" * 600)
    with pytest.raises(EngineError, match="needs a password to open"):
        open_data(base64.b64encode(write_cfb(root)).decode(), "secret.xlsx")
    old = Storage()
    old.set_stream("Workbook", b"\x09\x08" + b"\x00" * 600)
    with pytest.raises(EngineError, match="Legacy .xls"):
        open_data(base64.b64encode(write_cfb(old)).decode(), "old.xls")


def test_csv_of_two_cells_far_apart_is_refused_at_once():
    # #528: A1 and XFD1048576 would be 16 GiB of commas (and kept the server busy for ever).
    from gridcore.xlsx.csv_io import CsvTooLarge, write_csv

    s = new()
    s.run("cell.set", cell="A1", input="near")
    s.run("cell.set", cell="XFD1048576", input="far")
    t = time.time()
    with pytest.raises(EngineError, match="over the limit of 512 MiB"):
        s.run("file.saveBytes", format="csv")
    assert time.time() - t < 1
    # A sparse sheet is written quickly and exactly: only the filled cells are formatted.
    s = new()
    s.run("range.setValues", range="A1", values=[["a", "b,c", 'x"y'], [1, None, 3]])
    s.run("cell.set", cell="D5", input="=SEQUENCE(2)")
    s.run("cell.set", cell="B8", input="z")
    data = base64.b64decode(s.run("file.saveBytes", format="csv")["data"])
    assert data == b'\xef\xbb\xbfa,"b,c","x""y",\r\n1,,3,\r\n,,,\r\n,,,\r\n,,,1\r\n,,,2\r\n,,,\r\n,z,,\r\n'
    # The running size is checked too (long text in few cells).
    s = new()
    s.run("cell.set", cell="A1", input="x" * 3000)
    s.run("cell.set", cell="A2", input="y" * 3000)
    sh = s.doc().sheet
    with pytest.raises(CsvTooLarge):
        write_csv(sh, lambda rc: sh.value(rc) or "", limit=5000)
    assert len(write_csv(sh, lambda rc: sh.value(rc) or "", limit=7000)) > 6000


def test_saving_to_an_unknown_extension_is_refused():
    # #878
    s = new()
    s.run("cell.set", cell="A1", input="1")
    d = tempfile.mkdtemp()
    with pytest.raises(EngineError):
        s.run("file.saveAs", path=os.path.join(d, "report.cvs"))
    assert not os.listdir(d)
    s.run("file.saveAs", path=os.path.join(d, "report.xlsx"))
    assert os.listdir(d) == ["report.xlsx"]


# ---------------------------------------------------------------- commands


def test_a_monthly_series_counts_from_the_first_date():
    # #132: Jan 31, Feb 28, Mar 31, Apr 30.
    s = new()
    s.run("cell.set", cell="A1", input="2026-01-31")
    s.run("edit.fillSeries", range="A1:A4", type="date", unit="month")
    assert [s.run("cell.get", cell=f"A{i}")["text"] for i in range(1, 5)] == ["1/31/2026", "2/28/2026", "3/31/2026", "4/30/2026"]
    s.run("cell.set", cell="B1", input="2024-02-29")
    s.run("edit.fillSeries", range="B1:B3", type="date", unit="year")
    assert [s.run("cell.get", cell=f"B{i}")["text"] for i in range(1, 4)] == ["2/29/2024", "2/28/2025", "2/28/2026"]


def test_a_falling_growth_series_stops_at_its_stop_value():
    # b5a82554
    s = new()
    s.run("cell.set", cell="A1", input="100")
    s.run("edit.fillSeries", range="A1:A10", type="growth", step=0.5, stop=10)
    assert [val(s, f"A{i}") for i in range(1, 7)] == [100, 50, 25, 12.5, None, None]
    s.run("cell.set", cell="B1", input="1")
    s.run("edit.fillSeries", range="B1:B10", type="growth", step=2, stop=10)
    assert [val(s, f"B{i}") for i in range(1, 6)] == [1, 2, 4, 8, None]


def test_a_weekday_series_with_a_huge_step_is_refused_at_once():
    # #535: it used to keep the server busy for ever.
    s = new()
    s.run("cell.set", cell="A1", input="45000")
    t = time.time()
    with pytest.raises(EngineError):
        s.run("edit.fillSeries", range="A1:A3", type="date", unit="weekday", step=1e12)
    assert time.time() - t < 2 and val(s, "A2") is None
    s.run("cell.set", cell="B1", input="2026-10-16")  # a Friday
    s.run("edit.fillSeries", range="B1:B4", type="date", unit="weekday", step=1)
    assert [s.run("cell.get", cell=f"B{i}")["text"] for i in range(1, 5)] == ["10/16/2026", "10/19/2026", "10/20/2026", "10/21/2026"]


def test_autofill_times_quarters_ordinals_and_mixed_cells():
    # PR 483
    def fill(values, n):
        s = new()
        for i, v in enumerate(values):
            s.run("cell.set", cell=f"A{i + 1}", input=v)
        s.run("edit.autoFill", source=f"A1:A{len(values)}", target=f"A1:A{len(values) + n}")
        return [s.run("cell.get", cell=f"A{i + 1}")["text"] for i in range(len(values), len(values) + n)]

    assert fill(["9:00"], 2) == ["10:00", "11:00"]
    assert fill(["Qtr3"], 3) == ["Qtr4", "Qtr1", "Qtr2"]
    assert fill(["Quarter 3"], 2) == ["Quarter 4", "Quarter 1"]
    assert fill(["Q3"], 2) == ["Q4", "Q1"]
    assert fill(["1st Period"], 3) == ["2nd Period", "3rd Period", "4th Period"]
    assert fill(["11th"], 3) == ["12th", "13th", "14th"]
    assert fill(["Product 1", "On backorder"], 4) == ["Product 2", "On backorder", "Product 3", "On backorder"]
    assert fill(["Item 5"], 2) == ["Item 6", "Item 7"]  # unchanged
    # What Excel 16 for Mac does besides: the number a text starts with, else the one it ends with
    # (a number inside is copied); in a source without one pattern, neighbours of a kind continue
    # together and a number alone counts on by 1.
    assert fill(["a1b"], 2) == ["a1b", "a1b"]
    assert fill(["Item 5th"], 1) == ["Item 5th"]
    assert fill(["1 to 5"], 2) == ["2 to 5", "3 to 5"]
    assert fill(["1", "a"], 4) == ["2", "a", "3", "a"]
    assert fill(["1", "2", "x", "5"], 8) == ["3", "4", "x", "6", "5", "6", "x", "7"]
    assert fill(["Item 1", "Item 3", "x"], 6) == ["Item 5", "Item 7", "x", "Item 9", "Item 11", "x"]
    assert fill(["Mon", "Tue", "x"], 3) == ["Wed", "Thu", "x"]
    assert fill(["1", "", "a"], 3) == ["2", "", "a"]


def test_autofit_row_height_shows_hidden_rows():
    # PR 796, and what Excel 16 for Mac does: AutoFit Row Height shows the hidden rows it fits, empty
    # ones too, and SUBTOTAL follows; AutoFit Column Width shows a hidden column only when it holds
    # something.
    s = new()
    s.run("range.setValues", range="A1", values=[["H"], [1], [2], [3]])
    s.run("cell.set", cell="C1", input="=SUBTOTAL(109,A2:A4)")
    s.run("home.hideRows", rows="3:3")
    s.run("home.hideRows", rows="6:6")
    assert val(s, "C1") == 4
    s.run("home.autofitRowHeight", rows="2:6")
    assert not any(s.doc().sheet.is_row_hidden(r) for r in range(8)) and val(s, "C1") == 6
    s.run("cell.set", cell="E1", input="some text")
    s.run("home.hideColumns", cols="E:E")
    s.run("home.hideColumns", cols="F:F")
    s.run("home.autofitColumnWidth", cols="D:F")
    sh = s.doc().sheet
    assert not sh.col_info(4).hidden and sh.col_info(5).hidden


def test_removing_a_filter_without_criteria_keeps_hidden_rows():
    # #508
    s = new()
    s.run("range.setValues", range="A1", values=[["H"], [1], [2], [3], [4]])
    s.run("home.hideRows", rows="3:3")
    s.run("home.hideRows", rows="10:10")
    s.run("data.filter", range="A1:A5")
    s.run("data.filter")
    sh = s.doc().sheet
    assert [r + 1 for r in range(12) if sh.is_row_hidden(r)] == [3, 10]
    # A filter with criteria still shows the rows it hid when it goes.
    s.run("home.unhideRows", rows="1:12")
    s.run("data.filter", range="A1:A5")
    s.run("data.filterColumn", col="A", criterion={"kind": "values", "values": ["1"]})
    assert s.doc().sheet.is_row_hidden(2)
    s.run("data.filter")
    assert not any(s.doc().sheet.is_row_hidden(r) for r in range(6))


def test_committing_nothing_into_an_empty_cell_is_no_step():
    # #1855: Backspace, then Enter, is one undo step.
    s = new()
    s.run("cell.set", cell="A1", input="1")
    s.run("cell.set", cell="A2", input="2")
    s.run("edit.clearContents", range="A2")
    s.run("cell.set", cell="A2", input="")
    s.run("edit.undo")
    assert val(s, "A2") == 2 and val(s, "A1") == 1


def test_redo_selects_the_cells_of_the_step():
    # #1855
    s = new()
    s.run("selection.set", range="A1")
    s.run("cell.set", cell="A1", input="1")
    s.run("selection.set", range="D5")
    s.run("edit.undo")
    assert s.doc().selection.current().a1() == "A1"
    s.run("selection.set", range="F9")
    s.run("edit.redo")
    assert s.doc().selection.current().a1() == "A1" and val(s, "A1") == 1


def test_double_clicking_the_fill_handle_fills_down_to_the_adjacent_data():
    # #507
    s = new()
    s.run("range.setValues", range="A1", values=[[1, 10], [2, 20], [3, 30], [4, None], [5, None], [6, None]])
    s.run("selection.set", range="B1:B3")
    assert s.run("edit.autoFillDown") == {"range": "B1:B6"}
    assert [val(s, f"B{i}") for i in range(1, 8)] == [10, 20, 30, 40, 50, 60, None]
    assert s.doc().selection.current().a1() == "B1:B6"
    s.run("edit.undo")
    assert val(s, "B4") is None
    # It stops above a cell that holds something, and uses the right column when the left is empty.
    s = new()
    s.run("range.setValues", range="B1", values=[["=C1*2", 1], [None, 2], [None, 3], ["keep", 4], [None, 5]])
    s.run("selection.set", range="B1")
    s.run("edit.autoFillDown")
    assert [inp(s, f"B{i}") for i in range(1, 6)] == ["=C1*2", "=C2*2", "=C3*2", "keep", ""]
    s.run("selection.set", range="H1")
    assert s.run("edit.autoFillDown") == {"range": "H1"}
    # Both sides have data: as far as the longer run under the selection goes (Excel 16 for Mac);
    # a gap ends a run.
    for left, right, want in [(3, 6, "B1:B6"), (6, 3, "B1:B6"), (3, 2, "B1:B3")]:
        s = new()
        s.run("range.setValues", range="A1", values=[[i] for i in range(1, left + 1)])
        s.run("range.setValues", range="C1", values=[[i] for i in range(1, right + 1)])
        s.run("cell.set", cell="B1", input="=ROW()")
        s.run("selection.set", range="B1")
        assert s.run("edit.autoFillDown") == {"range": want}, (left, right)
    s = new()
    s.run("range.setValues", range="A1", values=[[1], [2], [None], [4], [5]])
    s.run("cell.set", cell="B1", input="=ROW()")
    s.run("selection.set", range="B1")
    assert s.run("edit.autoFillDown") == {"range": "B1:B2"}


def test_files_with_these_fixes_read_back_in_openpyxl():
    # External file links (#800), a header-and-totals table (#511), a column name with an edge space
    # (#732), functions by name (#902) and _xludf. (#726), read by another library.
    s = new()
    s.run("insert.hyperlink", cell="A1", target="Source.xlsx")
    s.run("insert.hyperlink", cell="A2", target="Sheet1!C3")
    s.run("range.setValues", range="C1", values=[["Name", "Qty "], ["a", 2], ["b", 3]])
    s.run("insert.table", range="C1:D3", name="Sales")
    s.run("table.totalRow", on=True, table="Sales")
    s.run("range.setValues", range="G1", values=[["H", "V"]])
    s.run("insert.table", range="G1:H1", name="Empty")
    s.run("table.totalRow", on=True, table="Empty")
    s.run("cell.set", cell="J1", input="=BYROW(D2:D3,SUM)")
    s.run("cell.set", cell="K1", input='=_xludf.CONCAT("a","b")')
    raw = base64.b64decode(s.run("file.saveBytes")["data"])
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        ws = openpyxl.load_workbook(io.BytesIO(raw)).active
    assert ws["A1"].hyperlink.target == "Source.xlsx" and not ws["A1"].hyperlink.location
    assert ws["A2"].hyperlink.location == "Sheet1!C3"
    assert dict(ws.tables.items()) == {"Sales": "C1:D4", "Empty": "G1:H3"}
    assert [c.name for c in ws.tables["Sales"].tableColumns] == ["Name", "Qty "] and ws["D1"].value == "Qty "
    assert ws["D4"].value == "=SUBTOTAL(109,Sales[[Qty ]])"
    assert getattr(ws["J1"].value, "text", ws["J1"].value) == "=_xlfn.BYROW(D2:D3,_xleta.SUM)"  # a spill: an array formula
    assert ws["K1"].value == '=_xludf.CONCAT("a","b")'


# ---------------------------------------------------------------- GridCraft's fixes of 2026-10-19 to 2026-10-26


def test_mod_is_the_exact_remainder_with_excels_quotient_limit():
    # c746eea8 made MOD an exact remainder with the divisor's sign, which is what Excel 16 for Mac
    # computes: MOD(-0.3,0.1) shows the rounding of 0.1. Excel also refuses a quotient of 1.1259E+12
    # or more (older Excels stopped at 2^27); GridCraft doesn't.
    assert ev("=MOD(-0.3,0.1)") == "2.77555756156289E-17"
    assert ev("=MOD(0.3,0.1)") == "0.1" and ev("=MOD(5,0.1)") == "0.0999999999999997"
    assert ev("=MOD(-3,2)") == "1" and ev("=MOD(3,-2)") == "-1" and ev("=MOD(2^39+0.75,0.5)") == "0.25"
    assert ev("=MOD(1.125E+12,1)") == "0" and ev("=MOD(2^40*1.01,1)") == "0.760009765625"
    for f in ("=MOD(1.13E+12,1)", "=MOD(5,1E-15)", "=MOD(1E+15,3)", "=MOD(2^53,10)", "=MOD(0.5,2^-42)", "=MOD(1,1E-14)"):
        assert ev(f) == "#NUM!", f
    assert ev("=MOD(3,0)") == "#DIV/0!"


def test_pmt_to_excels_last_bits_and_its_edges():
    # 1ce4781d: (1 + rate)^nper from log1p and expm1, which keeps the low bits of a small rate.
    # Excel 16 for Mac's answers; GridCraft's later 72c1d57d avoided the overflow of
    # PMT(1E-300,10,1E+12), which Excel reports as #NUM!.
    assert ev("=PMT(0.000001,360,100000)") == "-277.827919666642"
    assert ev("=PMT(1E-12,10,100)") == "-10.000000000055" and ev("=PMT(1E-10,100,1000)") == "-10.0000000505"
    assert ev("=PMT(0.05,1E+6,100)") == "-5" and ev("=PMT(1E-5,1E+8,100)") == "-0.001"
    assert ev("=PMT(1E-20,10,100)") == "-10" and ev("=PMT(1E-300,10,100)") == "-10"
    assert ev("=PMT(-0.999,5,100)") == "-9.99000000000009E-14"
    assert ev("=PMT(-0.5,60,1E+18)") == "-0.433680868994203"
    assert ev("=PMT(0.05,10,100,50,1)") == "-16.1197011855414" and ev("=PMT(0.05,-10,100)") == "7.95045749654567"
    for f in ("=PMT(-1,10,100)", "=PMT(-1.5,10,100)", "=PMT(1E-300,10,1E+12)", "=PMT(0.05,0,100)"):
        assert ev(f) == "#NUM!", f
    assert ev("=PPMT(1E-10,1,100,1000)") == "-9.9999999505"
    # FV and PV are unchanged: Excel's own sit further from the exact value than this form does.
    assert ev("=FV(1E-10,100,-10)") == "1000.00008274037" and ev("=PV(1E-10,100,-10)") == "1000.00007274037"


def test_accounting_negatives_read_one_pair_of_parentheses():
    # 35067fa9, checked in Excel 16 for Mac: $(1,234), $ (5), ($5), (5)%, (5%), (1 1/2) and € are
    # numbers; (-5), ((5)), $5%, (1/2), (10:30), (1/2/2024) and (5)$ stay text.
    s = new()
    currency, cents = "$#,##0_);[Red]($#,##0)", "$#,##0.00_);[Red]($#,##0.00)"
    numbers = {
        "$(1,234)": (-1234.0, currency), "$ (5)": (-5.0, currency), "($5)": (-5.0, currency), "( $ 1,234 )": (-1234.0, currency),
        "$(1,234.)": (-1234.0, cents), "$(.5)": (-0.5, cents), "$(0)": (0.0, currency), "(5)%": (-0.05, "0%"), "(5%)": (-0.05, "0%"),
        "( 5 ) %": (-0.05, "0%"), "(1,234)%": (-12.34, "0%"), "(5.)%": (-0.05, "0.00%"), "(1 1/2)": (-1.5, "# ?/?"),
        "(1e-3)": (-0.001, "0.00E+00"), "(5)": (-5.0, "General"), "( 5 )": (-5.0, "General"), "(.5)": (-0.5, "General"), "€5": (5.0, "General"),
    }
    for text, (value, fmt) in numbers.items():
        s.run("home.numberFormat", range="A1", code="General")
        s.run("cell.set", cell="A1", input=text)
        got = s.run("cell.get", cell="A1")
        assert (got["value"], got["format"]) == (value, fmt), (text, got)
    for text in ["(-5)", "(+5)", "((5))", "$5%", "$(5)%", "(5)%%", "($-5)", "(1/2)", "(10:30)", "(1/2/2024)", "(5)$", "(5)(5)", "$(5", "5)", "(1,23)", "$(1/2)", "(TRUE)"]:
        s.run("home.numberFormat", range="A1", code="General")
        s.run("cell.set", cell="A1", input=text)
        assert s.run("cell.get", cell="A1")["value"] == text, text
    # VALUE and arithmetic read the same way.
    assert ev('=VALUE("$(1,234)")') == "-1234" and ev('="(5)"*1') == "-5" and ev('=VALUE("(1 1/2)")') == "-1.5"
    for f in ('=VALUE("((5))")', '=VALUE("(-5)")', '=VALUE("$5%")', '=VALUE("(1/2)")', '=VALUE("(10:30)")', '=VALUE("$(5)%")'):
        assert ev(f) == "#VALUE!", f
    # #2163: digits too many for a finite number don't make one (∞/∞ was NaN).
    assert ev('=VALUE("1 "&REPT("9",400)&"/"&REPT("9",400))') == "#VALUE!"
    assert ev('=("1"&REPT("0",400)&"/"&REPT("9",400))+0') == "#VALUE!"


def test_the_text_of_a_number_is_written_out_up_to_20_characters():
    # f29ac7d9 made 999999999999999.9 "1E+15"; Excel 16 for Mac writes a number out in full while
    # that takes at most 20 characters (the sign not counted) and in scientific notation past that.
    cases = [
        ('=(999999999999999+0.9)&""', "1000000000000000"), ('=2^63&""', "9223372036854780000"), ('=1E-18&""', "0.000000000000000001"),
        ('=1E-19&""', "1E-19"), ('=1E+20&""', "1E+20"), ('=(1E+16/3)&""', "3333333333333330"), ('=(1/3*1E-5)&""', "3.33333333333333E-06"),
        ('=(2/3*1E-3)&""', "0.000666666666666667"), ('=(1/7*1E-3)&""', "0.000142857142857143"), ('=(1/7*1E-4)&""', "1.42857142857143E-05"),
        ('=-(1E+19/7)&""', "-1428571428571430000"), ('=99999999999999900000&""', "99999999999999900000"), ('=(1E+15+0.5)&""', "1000000000000000"),
        ('=1E-10&""', "0.0000000001"), ('=(1/3)&""', "0.333333333333333"), ('=(2^53)&""', "9007199254740990"),
    ]
    for f, want in cases:
        assert ev(f) == want, f
    # 9cbb674c: a formula's constants keep their short form, and never round past the largest number.
    s = new()
    s.run("cell.set", cell="A1", input="=1E-10+0")
    assert inp(s, "A1") == "=1E-10+0"
    s.run("cell.set", cell="A2", input="=1.7976931348623157E+308")
    assert inp(s, "A2") == "=1.7976931348623157E+308" and val(s, "A2") == 1.7976931348623157e308


def test_isoweeknum_at_the_start_of_the_calendar():
    # 86064077: the week of serial 1 belongs to 1899 (Excel 16 for Mac: ISOWEEKNUM(1) is 52, not 0).
    assert [ev(f"=ISOWEEKNUM({x})") for x in (0, 1, 2, 7, 8, 9)] == ["52", "52", "1", "1", "1", "2"]
    assert ev("=WEEKNUM(1,21)") == "52" and ev("=WEEKNUM(0,21)") == "52" and ev("=WEEKNUM(8,21)") == "1"
    assert ev("=ISOWEEKNUM(-1)") == "#NUM!"
    # A 1904 workbook: January 1 to 3, 1904 are in week 53 of 1903.
    wb = wb1904()
    assert [ev(f"=ISOWEEKNUM({x})", wb) for x in (0, 1, 2)] == ["53", "53", "53"]


def test_imsqrt_follows_excels_polar_form():
    # 92d2b989 removed the stray real part; Excel 16 for Mac keeps it (the polar form), counts a
    # squared magnitude below the smallest normal number as 0 and overflows to #NUM!.
    assert ev('=IMSQRT("-4")') == "1.22464679914735E-16+2i" and ev('=IMSQRT("-1")') == "6.12323399573677E-17+i"
    assert ev('=IMSQRT("-2")') == "8.65956056235493E-17+1.4142135623731i"
    assert ev('=IMSQRT("-2i")') == "1-i" and ev('=IMSQRT("3+4i")') == "2+i" and ev('=IMSQRT("4")') == "2"
    assert ev('=IMSQRT("i")') == "0.707106781186548+0.707106781186547i"
    assert ev('=IMSQRT("-1E150")') == "6.12323399573677E+58+1E+75i"
    assert ev('=IMSQRT("1E154")') == "1E+77" and ev('=IMSQRT("1E155")') == "#NUM!"
    assert ev('=IMSQRT("1E-154")') == "0" and ev('=IMSQRT("1E-300")') == "0"


def test_percentrank_significance_past_the_number_range():
    # 14415f41 returned the rank untruncated; Excel 16 for Mac gives #DIV/0! past 308.
    assert ev("=PERCENTRANK({1,2,3,4},2.5,308)") == "0.5" and ev("=PERCENTRANK({1,2,3,4},2,308)") == "0.333333333333333"
    for f in ("=PERCENTRANK({1,2,3,4},2.5,309)", "=PERCENTRANK({1,2,3,4},2,400)", "=PERCENTRANK.EXC({1,2,3,4},2.5,1E+10)"):
        assert ev(f) == "#DIV/0!", f
    assert ev("=PERCENTRANK({1,2,3,4},2.5,0)") == "#NUM!"


def test_sheet_names_are_quoted_as_excel_quotes_them():
    # #1458 and #2473, checked in Excel 16 for Mac: names that start with a dot or read as R1C1
    # (R, C, RC, R1C, RC1, r2c3) or as a cell are quoted; past the sheet's edge (XFE1, AB1048577)
    # they are not.
    s = new()
    for name, quoted in [(".foo", True), ("RC", True), ("R1C", True), ("r2c3", True), ("CR1", True), ("AB1048576", True), ("TRUE", True),
                         ("A1A", False), ("Sheet.1", False), ("RR", False), ("XFE1", False), ("AB1048577", False), ("_x", False), ("Ärger", False)]:
        s.run("home.insertSheet", name=name)
        s.run("sheet.activate", sheet="Sheet1")
        s.run("cell.set", cell="A1", input=f"='{name}'!A1")
        assert inp(s, "A1") == (f"='{name}'!A1" if quoted else f"={name}!A1"), name
    # ADDRESS follows the same rule; a [Book]Sheet prefix is judged with its sheet part's letters.
    for sheet, want in [("RC", "'RC'!$A$1"), (".x", "'.x'!$A$1"), ("TRUE", "'TRUE'!$A$1"), ("XFE1", "XFE1!$A$1"), ("[Book1]Sheet1", "[Book1]Sheet1!$A$1"),
                        ("[Book1]A1", "[Book1]A1!$A$1"), ("[Book 1]Sheet1", "'[Book 1]Sheet1'!$A$1"), ("My Sheet", "'My Sheet'!$A$1")]:
        assert ev(f'=ADDRESS(1,1,1,TRUE,"{sheet}")') == want, sheet
    # #2473: a range whose second end repeats the sheet (in any case) is accepted, as typed.
    s.run("range.setValues", range="B1", values=[[1, 2], [3, 4]])
    s.run("cell.set", cell="D1", input="=SUM(Sheet1!B1:SHEET1!C2)")
    assert val(s, "D1") == 10 and inp(s, "D1") == "=SUM(Sheet1!B1:SHEET1!C2)"


def test_let_and_lambda_names_fold_case_as_text_does():
    # ac371be5 and a74e82ac: dług is DŁUG, ß is ẞ, ς is σ (Excel 16: ="ς"="σ" is TRUE); ı and İ
    # are letters of their own.
    assert ev("=LET(dług,100,DŁUG*2)") == "200" and ev("=LET(ς,5,σ)") == "5" and ev("=LET(ß,1,ẞ)") == "1"
    assert ev("=LET(i,1,ı,2,i)") == "1" and ev("=LET(İ,1,i)") == "#NAME?"
    assert ev('="ς"="σ"') == "TRUE" and ev('="ı"="i"') == "FALSE"


def test_a_dynamic_array_of_one_value_still_anchors_a_spill_range():
    # #1491: C1# gave #REF! while the FILTER in C1 matched one row, so ROWS, SUM and COUNT over
    # it failed the moment a list had one entry.
    s = new()
    s.run("range.setValues", range="A1", values=[[1, "x"], [2, "y"], [3, "z"]])
    s.run("cell.set", cell="C1", input='=FILTER(A1:A3,B1:B3="x")')
    s.run("cell.set", cell="E1", input="=SORT(A1:A1)")
    s.run("cell.set", cell="F1", input="=1+1")
    assert (cell(s, "=ROWS(C1#)", "D5"), cell(s, "=SUM(C1#)", "D6"), cell(s, "=COUNT(C1#)", "D7")) == (1, 1, 1)
    assert cell(s, "=ROWS(E1#)", "D8") == 1
    assert display(cell(s, "=ROWS(F1#)", "D9")) == "#REF!"  # a scalar formula is no anchor
    assert s.doc().wb.sheet(0).spill_ranges[(0, 2)] == RangeRef(0, 2, 0, 2)


def test_cell_filename_names_the_workbooks_file():
    # #1489: it was "" for every workbook. Excel gives <folder>\[<file>]<sheet>, and "" before the
    # first save; a workbook that arrived as bytes knows no folder.
    s = new()
    assert cell(s, '=CELL("filename",A1)') == ""
    s2 = roundtrip(s, "Budget.xlsx")
    assert cell(s2, '=CELL("filename",A1)') == "[Budget.xlsx]Sheet1"
    assert cell(s2, '=MID(CELL("filename",A1),FIND("]",CELL("filename",A1))+1,31)') == "Sheet1"


def test_hostile_formulas_are_refused_at_once():
    # #2182: the broadcast of two arrays can be 17 billion cells; #2175: SUBTOTAL over A:XFD with a
    # value in XFD1048576; #2415: star churn in SEARCH; #2177: a regression with 16,384 predictors.
    # Each used to keep the server busy for minutes or for ever.
    s = new()
    s.run("cell.set", cell="XFD1048576", input="1")
    t = time.time()
    for f in ("=ROWS(SEQUENCE(1048576)*SEQUENCE(1,16384))", "=ROWS(IF(SEQUENCE(1048576)>0,SEQUENCE(1,16384),0))", "=SUBTOTAL(9,E:XFD)",
              "=AGGREGATE(9,0,E:XFD)", "=INDEX(LINEST(SEQUENCE(3),SEQUENCE(3,16384)),1,1)"):
        assert display(cell(s, f, "A1")) == "#NUM!", f
    assert display(cell(s, '=SEARCH("*a*a*a*a*a*a*a*a*a*a*a*a*b",REPT("a",30000))', "A2")) == "#VALUE!"
    assert time.time() - t < 20
    # Honest work is untouched.
    # Excel 16 for Mac: a leading star absorbs the text before the match, so the match starts at
    # start_num; a start from the end of the text on is #VALUE!, an empty find_text included.
    assert ev('=SEARCH("*a*b","xaab")') == "1" and ev('=SEARCH("*b","aab",2)') == "2" and ev('=SEARCH("*b","aab",3)') == "3"
    assert ev('=SEARCH("a*","xxa")') == "3" and ev('=SEARCH("?b","aab")') == "2" and ev('=SEARCH("*~*b","a*b")') == "1"
    assert ev('=SEARCH("?*zz",REPT("a",30000)&"zz")') == "1"
    for f in ('=SEARCH("*b","aab",4)', '=SEARCH("","abc",4)', '=SEARCH("*","abc",4)', '=SEARCH("*?c","abc",3)', '=SEARCH("*q","xaab")', '=SEARCH("b*","abc",3)', '=SEARCH("*","")'):
        assert ev(f) == "#VALUE!", f
    assert ev('=MATCH("*b",{"xab","b"},0)') == "1"
    assert ev('=SUM(SEQUENCE(2000)*SEQUENCE(1,100))') == "10105050000"


def test_an_autofilter_grows_when_a_value_is_entered_in_the_row_below_it():
    # #2342 and #2367
    s = new()
    s.run("range.setValues", range="A1", values=[["Name", "Qty"], ["d", 4], ["b", 2]])
    s.run("data.filter", range="A1:B3")

    def rng():
        return s.doc().wb.sheet(0).autofilter.range.a1()

    s.run("cell.set", cell="A4", input="e")
    assert rng() == "A1:B4"
    s.run("range.setValues", range="A5", values=[["f", 6]])
    assert rng() == "A1:B5"
    s.run("cell.set", cell="A7", input="x")  # after a blank row
    s.run("cell.set", cell="C6", input="x")  # outside its columns
    s.run("range.setValues", range="A6", values=[[""]])  # empty text is no entry
    assert rng() == "A1:B5"
    s.run("cell.set", cell="B6", input="=1+1")
    assert rng() == "A1:B6"
    s.run("range.fill", range="A7:B7", input="g")
    assert rng() == "A1:B7"


def test_sparklines_follow_rows_and_go_with_a_deleted_row():
    # 80e20411: a sparkline whose row was deleted stayed, pointing at the next row's cell (and
    # sparklines never followed inserted rows here).
    s = new()
    s.run("range.setValues", range="A1", values=[[1, 2, 3], [4, 5, 6]])
    s.run("insert.sparkline", data="A1:C1", location="E1")
    s.run("insert.sparkline", data="A2:C2", location="E2")
    s.run("home.insertRows", rows="1:1")
    sp = sorted((x.row, x.col, x.source) for x in s.doc().wb.sheet(0).sparklines)
    assert sp == [(1, 4, "Sheet1!$A$2:$C$2"), (2, 4, "Sheet1!$A$3:$C$3")], sp
    s.run("home.deleteRows", rows="2:2")
    sp = [(x.row, x.col, x.source) for x in s.doc().wb.sheet(0).sparklines]
    assert sp == [(1, 4, "Sheet1!$A$2:$C$2")], sp


def test_a_conditional_format_threshold_formula_must_parse():
    # #1999: a formula that doesn't parse must never reach the workbook — Excel refuses to open a
    # file carrying it, even in repair mode.
    s = new()
    with pytest.raises(EngineError):
        s.run("home.conditionalFormat", range="A1:A5", rule={"type": "dataBar", "min": {"kind": "formula", "value": "=1+("}})
    s.run("home.conditionalFormat", range="A1:A5", rule={"type": "dataBar", "min": {"kind": "formula", "value": "=$B$1"}})
    assert s.doc().wb.sheet(0).cond_formats[0].rule.bar_min.value == "$B$1"


def test_row_auto_fit_counts_line_breaks_only_in_wrapping_cells_and_wraps_whole_words():
    # 4faa3ff0 and #1969
    from gridcore.engine.layout import approx_text_width, wrapped_line_count

    s = new()
    s.run("cell.set", cell="A1", input="one\ntwo\nthree")  # a typed break turns wrap on and grows the row
    sh = s.doc().wb.sheet(0)
    info = sh.rows.get(0)
    assert info is not None and info.size is not None and info.size > sh.default_row_height
    s.run("home.wrapText", on=False, range="A1")
    s.run("home.autofitRowHeight", rows="1:1")
    info = s.doc().wb.sheet(0).rows.get(0)
    assert info is None or info.size is None
    # Twelve five-letter words in a column that holds two of them and a bit: whole words wrap, so six
    # lines, where the width alone says five.
    w = approx_text_width("abcde", 11.0, False)
    sp = approx_text_width(" ", 11.0, False)
    line = " ".join(["abcde"] * 12)
    assert wrapped_line_count(line, 2.9 * w + 2 * sp, 11.0, False) == 6
    assert math.ceil(approx_text_width(line, 11.0, False) / (2.9 * w + 2 * sp)) == 5


def test_fill_days_weekdays_months_and_years_keep_the_time_of_day():
    # 00edfd06
    s = new()
    s.run("cell.set", cell="A1", input="1/31/2026 9:30")
    s.run("edit.autoFill", source="A1", target="A1:A4", mode="months")
    assert [cell(s, f'=TEXT(A{r},"yyyy-mm-dd hh:mm")', "J1") for r in (2, 3, 4)] == ["2026-02-28 09:30", "2026-03-31 09:30", "2026-04-30 09:30"]
    s.run("cell.set", cell="B1", input="1/2/2026")  # a Friday
    s.run("edit.autoFill", source="B1", target="B1:B4", mode="weekdays")
    assert [cell(s, f'=TEXT(B{r},"ddd d")', "J1") for r in (2, 3, 4)] == ["Mon 5", "Tue 6", "Wed 7"]
    s.run("range.setValues", range="D1", values=[["1/1/2026"], ["1/3/2026"]])
    s.run("edit.autoFill", source="D1:D2", target="D1:D4", mode="years")
    assert [cell(s, f'=TEXT(D{r},"yyyy-m-d")', "J1") for r in (3, 4)] == ["2027-1-3", "2028-1-3"]
    s.run("range.setValues", range="E1", values=[["1/1/2026"], ["1/4/2026"]])
    s.run("edit.autoFill", source="E1:E2", target="E1:E4", mode="days")
    assert [cell(s, f'=TEXT(E{r},"m/d")', "J1") for r in (3, 4)] == ["1/7", "1/10"]
    # A source that isn't dates fills as a plain drag does.
    s.run("cell.set", cell="F1", input="5")
    s.run("edit.autoFill", source="F1", target="F1:F3", mode="months")
    assert (val(s, "F2"), val(s, "F3")) == (5, 5)


def test_editing_a_note_keeps_its_box_and_who_wrote_it():
    # #1608: a note's box was written at a fixed place and size; 3e6eac42: Edit Note replaced the
    # whole note.
    from dataclasses import replace

    s = new()
    s.run("review.newNote", cell="B2", text="first", author="Ann")
    sh = s.doc().wb.sheet(0)
    sh.comments[(1, 1)] = replace(sh.comments[(1, 1)], box=(2, 10, 1, 5, 5, 20, 6, 8, 150.0, 80.0))
    vml = part(s, "xl/drawings/vmlDrawing1.vml")
    assert "<x:Anchor>3, 10, 2, 5, 6, 20, 7, 8</x:Anchor>" in vml and "width:150pt;height:80pt" in vml
    s2 = roundtrip(s)
    assert s2.doc().wb.sheet(0).comments[(1, 1)].box == (2, 10, 1, 5, 5, 20, 6, 8, 150.0, 80.0)
    s2.run("review.editNote", cell="B2", text="second")
    c = s2.doc().wb.sheet(0).comments[(1, 1)]
    assert (c.text, c.author, c.box) == ("second", "Ann", (2, 10, 1, 5, 5, 20, 6, 8, 150.0, 80.0))
    # A note made here keeps the default box, which is not stored.
    s.run("review.newNote", cell="C3", text="plain")
    assert s.doc().wb.sheet(0).comments[(2, 2)].box is None
    assert "<x:Anchor>3, 15, 2, 2, 5, 15, 6, 16</x:Anchor>" in part(s, "xl/drawings/vmlDrawing1.vml")


def _rewritten(s, name, fn):
    """The workbook ``s`` saves, with the part ``name`` changed by ``fn``, as base64."""
    raw = base64.b64decode(s.run("file.saveBytes")["data"])
    src = zipfile.ZipFile(io.BytesIO(raw))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == name:
                data = fn(data.decode()).encode()
            dst.writestr(item, data)
    return base64.b64encode(buf.getvalue()).decode()


def test_sheets_keep_their_sheet_ids():
    # 85e8f790: a rewritten workbook renumbered sheetIds from 1, so anything that recognises a
    # sheet by its id lost it.
    s = new()
    s.run("home.insertSheet", name="B")
    s.run("home.insertSheet", name="C")
    data = _rewritten(s, "xl/workbook.xml", lambda x: x.replace('sheetId="2"', 'sheetId="7"').replace('sheetId="3"', 'sheetId="9"'))
    s2 = open_data(data, "ids.xlsx")
    before = {x.name: x.sheet_id for x in s2.doc().wb.sheets}
    assert sorted(before.values()) == [1, 7, 9] and before["B"] == 7
    s2.run("home.deleteSheet", sheet="B")
    s2.run("home.insertSheet", name="D")
    ids = dict(re.findall(r'<sheet name="([^"]+)" sheetId="(\d+)"', part(s2, "xl/workbook.xml")))
    # The sheets keep their ids; the new one takes the first id no sheet asks for.
    assert ids == {"Sheet1": str(before["Sheet1"]), "C": str(before["C"]), "D": "2"}, ids


def test_fonts_keep_their_charset_family_and_theme_link():
    # 24b0b5d1, 6760e006 and #2517: the charset (178 = Arabic), the family class and the link to
    # the theme's body font survive a save, on every font that had it.
    from xml.etree import ElementTree as ET

    from gridcore.model.style import Font, Theme
    from gridcore.xlsx.styles import _font_xml, read_font
    from gridcore.xlsx.tables import INDEXED

    ns = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    bold = read_font(ET.fromstring(f'<font {ns}><b/><sz val="11"/><name val="Calibri"/><family val="2"/><scheme val="minor"/></font>'), list(INDEXED))
    arabic = read_font(ET.fromstring(f'<font {ns}><sz val="11"/><name val="Arial"/><family val="3"/><charset val="178"/></font>'), list(INDEXED))
    assert (bold.bold, bold.scheme, bold.family, bold.charset) == (True, "minor", None, None)
    assert (arabic.charset, arabic.family, arabic.scheme) == (178, 3, "")
    assert _font_xml(bold, False, Theme()).endswith('<name val="Calibri"/><family val="2"/><scheme val="minor"/></font>')
    assert _font_xml(arabic, False, Theme()).endswith('<name val="Arial"/><family val="3"/><charset val="178"/></font>')
    # Calibri chosen by hand (no link in the file) gets no scheme; a new font has the link; a font
    # given another face than the theme's keeps none.
    from dataclasses import replace

    assert "scheme" not in _font_xml(replace(arabic, name="Calibri"), False, Theme())
    assert _font_xml(Font(), False, Theme()).endswith('<scheme val="minor"/></font>')
    assert "scheme" not in _font_xml(replace(Font(), name="Arial"), False, Theme())
    # Through a file: the bold font derived from Normal keeps the link.
    s = new()
    s.run("cell.set", cell="A1", input="x")
    s.run("home.bold", range="A1", on=True)
    assert part(s, "xl/styles.xml").count('<scheme val="minor"/>') == 2
    s.run("home.fontName", range="A1", name="Arial")
    assert '<name val="Arial"/><family val="2"/></font>' in part(s, "xl/styles.xml")


def test_accounting_formats_with_a_currency_are_written_with_their_code():
    # b389d906: ids 42 and 44 are resolved by the reader's locale (LibreOffice showed "- 0.14" where
    # US Excel shows "$ (0.14)"), so they are written as custom formats like the currency formats 5–8.
    from gridcore.xlsx.tables import builtin_format

    s = new()
    s.run("cell.set", cell="A1", input="-0.14")
    s.run("home.numberFormat", range="A1", code=builtin_format(44))
    s.run("cell.set", cell="A2", input="-0.14")
    s.run("home.numberFormat", range="A2", code=builtin_format(43))
    styles = part(s, "xl/styles.xml")
    assert 'numFmtId="44"' not in styles and 'numFmtId="43"' in styles
    assert re.search(r'<numFmt numFmtId="1\d\d" formatCode="[^"]*&quot;\$&quot;', styles), styles
    s2 = roundtrip(s)
    assert s2.run("cell.get", cell="A1")["format"] == builtin_format(44)


def test_phonetic_runs_survive_open_and_save():
    # 40803e6a and e42ebf93: a Japanese workbook's furigana (<rPh> runs and <phoneticPr>) were
    # deleted by an open and save.
    s = new()
    s.run("range.setValues", range="A1", values=[["東京"], ["大阪"], ["東京"]])
    runs = '<si><t>東京</t><rPh sb="0" eb="2"><t>トウキョウ</t></rPh><phoneticPr fontId="1" type="Hiragana" alignment="center"/></si>'
    damaged = '<si><t>大阪</t><rPh sb="5" eb="9"><t>x</t></rPh><rPh sb="1" eb="1"><t>y</t></rPh><phoneticPr fontId="0" type="bogus"/></si>'
    data = _rewritten(s, "xl/sharedStrings.xml", lambda x: x.replace("<si><t>東京</t></si>", runs).replace("<si><t>大阪</t></si>", damaged))
    s2 = open_data(data, "furigana.xlsx")
    assert s2.doc().wb.phonetics == {"東京": ([(0, 2, "トウキョウ")], {"fontId": 1, "type": "Hiragana", "alignment": "center"}),
                                    "大阪": ([], {"fontId": 0, "type": None, "alignment": None})}
    sst = part(s2, "xl/sharedStrings.xml")
    assert runs in sst and '<si><t>大阪</t><phoneticPr fontId="0"/></si>' in sst
    assert cell(s2, "=A1&A2", "C1") == "東京大阪"


def test_the_themes_east_asian_and_per_script_typefaces_survive():
    # 2c6ffe48 and 189eb7ae: a Japanese workbook's theme fonts were deleted by an open and save.
    s = new()
    major = '<a:ea typeface="游ゴシック Light"/><a:cs typeface="Arial"/><a:font script="Jpan" typeface="游ゴシック Light"/></a:majorFont>'
    data = _rewritten(s, "xl/theme/theme1.xml", lambda x: x.replace('<a:ea typeface=""/><a:cs typeface=""/></a:majorFont>', major, 1))
    s2 = open_data(data, "theme.xlsx")
    t = s2.doc().wb.theme
    assert (t.major_font_ea, t.major_font_cs, t.major_font_scripts, t.minor_font_ea) == ("游ゴシック Light", "Arial", (("Jpan", "游ゴシック Light"),), "")
    assert major in part(s2, "xl/theme/theme1.xml")


def test_a_100_percent_stacked_area_chart_opens_and_saves_as_such():
    # 0eb2f1f6: an area chart with percentStacked grouping opened as a stacked area.
    s = new()
    s.run("range.setValues", range="A1", values=[["x", "a", "b"], [1, 2, 3], [2, 3, 4]])
    s.run("selection.set", range="A1:C3")
    s.run("insert.chart", kind="areaStacked100")
    assert '<c:areaChart><c:grouping val="percentStacked"/>' in part(s, "xl/charts/chart1.xml")
    s2 = roundtrip(s)
    assert s2.doc().wb.sheet(0).charts[0].kind == "areaStacked100"


def test_format_as_table_inside_a_table_restyles_it():
    # 09f78cb3
    s = new()
    s.run("range.setValues", range="A1", values=[["h", "k"], [1, 2], [3, 4]])
    s.run("insert.table", range="A1:B3", name="T")
    assert s.run("insert.table", range="A2:B2", style="TableStyleLight9") == {"name": "T", "range": "A1:B3", "restyled": True}
    assert s.doc().wb.sheet(0).tables[0].style == "TableStyleLight9"
    with pytest.raises(EngineError):
        s.run("insert.table", range="A2:B2")  # Insert › Table still refuses to overlap


def test_autosum_on_a_whole_row_totals_it_after_the_data():
    # 802c7823
    s = new()
    s.run("range.setValues", range="B2", values=[[1, 2, 3, 4]])
    s.run("selection.set", range="2:2")
    assert s.run("formulas.autoSum") == {"cells": ["F2"]}
    assert inp(s, "F2") == "=SUM(A2:E2)" and val(s, "F2") == 10


def test_an_explicit_range_keeps_its_blanks_and_a_padded_stack_is_capped():
    # #1010: a range of more than 16,384 cells was trimmed to the used range, so COUNTBLANK(A1:A20000)
    # was 1 on an empty sheet (Excel 16: 20000); only whole columns and rows, and ranges past a
    # million cells, trim now. #536: HSTACK/VSTACK pad to the tallest/widest array, which is what
    # must fit (HSTACK(SEQUENCE(1048576),SEQUENCE(1,16384)) used to run for ever).
    s = new()
    assert (cell(s, "=COUNTBLANK(A1:A20000)"), cell(s, '=SUMPRODUCT(--(B1:B20000=""))'), cell(s, '=COUNTIF(C1:C20000,"")')) == (20000, 20000, 20000)
    t = time.time()
    for f in ("=ROWS(HSTACK(SEQUENCE(1048576),SEQUENCE(1,16384)))", "=ROWS(VSTACK(SEQUENCE(1,16384),SEQUENCE(1048576)))"):
        assert display(cell(s, f)) == "#NUM!", f
    assert time.time() - t < 10
    assert cell(s, "=ROWS(HSTACK(SEQUENCE(3),SEQUENCE(2,2)))") == 3 and cell(s, "=COLUMNS(VSTACK(SEQUENCE(1,3),SEQUENCE(2)))") == 3
