"""The workbook list (``/``): who is editing which workbook, a new workbook from the list and its
title, search and filters, uploading an .xlsx, downloading one, and deleting a workbook while
someone has it open (they are told)."""

import sys
import time
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from _lib.grid import Grid, run_people  # noqa: E402


def wait_until(fn, what: str, timeout: float = 10.0):
    end = time.time() + timeout
    last = None
    while time.time() < end:
        last = fn()
        if last:
            return last
        time.sleep(0.15)
    raise AssertionError(f"timed out waiting for {what} (last: {last!r})")


def go_home(g: Grid) -> None:
    g.page.click("#home-link")
    g.page.locator("#manager").wait_for()
    g.wait_idle(0.6)


def rows(g: Grid):
    return g.page.locator(".rg-book-row")


def row(g: Grid, title: str):
    return g.page.locator(".rg-book-row", has=g.page.locator(".rg-book-link", has_text=title))


def body(grids) -> None:
    me, other = grids
    page = me.page
    first = me.book

    # The list: the workbook the other person has open, being edited.
    go_home(me)
    assert rows(me).count() == 1, rows(me).count()
    wait_until(lambda: "1 editing" in row(me, "Untitled workbook").inner_text(), "1 editing")

    # A new workbook from the list, renamed.
    page.click("#new-book")
    page.click("#new-blank")
    me.wait_title("Untitled workbook")
    assert "/book/" in page.url and first not in page.url, page.url
    t = page.locator("#doc-title")
    t.click()
    t.fill("Quarterly plan")
    t.press("Enter")
    me.type_in("A1", "Q1")
    me.wait_idle(2.0)  # saved a moment after the change
    go_home(me)
    row(me, "Quarterly plan").wait_for()
    assert rows(me).count() == 2

    # Search and filters.
    search = page.locator("#book-search")
    search.fill("quarterly")
    wait_until(lambda: rows(me).count() == 1, "one match")
    search.fill("zzz-nothing")
    page.get_by_text("No workbooks match these filters").wait_for()
    page.get_by_role("button", name="Reset filters").first.click()
    wait_until(lambda: rows(me).count() == 2, "both again")
    page.click("#filter-live")
    wait_until(lambda: rows(me).count() == 1, "only the one being edited")
    assert row(me, "Untitled workbook").count() == 1
    page.click("#filter-live")
    wait_until(lambda: rows(me).count() == 2, "both again")

    # Upload an .xlsx: it opens as a new shared workbook, titled after the file.
    src = me.out / "Upload me.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Item"
    ws["B1"] = 3
    ws["B2"] = "=B1*14"
    wb.save(src)
    page.locator('#manager input[type="file"]').set_input_files(str(src))
    me.wait_title("Upload me")
    me.wait_text("B2", "42")
    go_home(me)
    row(me, "Upload me").wait_for()
    assert rows(me).count() == 3

    # Download one from the list.
    with page.expect_download(timeout=20000) as dl:
        row(me, "Quarterly plan").get_by_title("Download .xlsx").click()
    out = me.out / "downloaded.xlsx"
    dl.value.save_as(out)
    assert dl.value.suggested_filename == "Quarterly plan.xlsx", dl.value.suggested_filename
    assert openpyxl.load_workbook(out)["Sheet1"]["A1"].value == "Q1"
    me.shot("list")

    # Delete the workbook the other person is in: they are told.
    row(me, "Untitled workbook").locator(".rg-delete").click()
    page.locator(".rg-delete-ok").click()
    wait_until(lambda: rows(me).count() == 2, "deleted from the list")
    other.page.locator("#book-gone").wait_for(state="visible", timeout=4000)
    other.shot("gone")


if __name__ == "__main__":
    sys.exit(run_people("workbook_manager", body, count=2))
