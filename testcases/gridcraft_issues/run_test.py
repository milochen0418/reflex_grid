"""The browser side of bugs GridCraft (the Rust spreadsheet gridcore follows) fixed: a header click in
formula point mode, Enter pasting a copy, typing after the name box, Ctrl+Shift+Arrow revealing the
moving end, the last row reached and drawn, a new chart selected so Delete removes it,
double-clicking the fill handle, a validation's input message under its open list, Center Across
Selection, Shrink to Fit and a merged area's borders, links that may not run script, and panning the
grid with a finger."""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402
from openpyxl.styles import Alignment, Border, Side  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from _lib.grid import MOD, Grid, run  # noqa: E402


def header(g: Grid, kind: str, label: str):
    """The center of a column (``cols``) or row (``rows``) header."""
    heads = g.page.locator(".rg-colhead" if kind == "cols" else ".rg-rowhead")
    for i in range(heads.count()):
        el = heads.nth(i)
        if el.inner_text().strip() == label:
            b = el.bounding_box()
            return b["x"] + b["width"] / 2, b["y"] + b["height"] / 2
    raise AssertionError(f"no {kind} header {label}")


def editor(g: Grid) -> str:
    return g.page.locator("#rg-editor").input_value()


def toast(g: Grid, text: str) -> None:
    g.page.locator("[data-sonner-toast]", has_text=text).first.wait_for(state="visible", timeout=6000)


def body(g: Grid) -> None:
    page = g.page
    for r in range(1, 5):
        g.type_in(f"C{r}", str(r))
        g.type_in(f"D{r}", str(r * 10))

    # GridCraft #691: in point mode a column header inserts C:C, a row header 4:4; a drag across
    # headers extends the reference.
    g.click("A1")
    page.keyboard.type("=SUM(")
    page.mouse.click(*header(g, "cols", "C"))
    g.wait_idle(0.2)
    assert editor(g) == "=SUM(C:C", editor(g)
    page.mouse.move(*header(g, "cols", "C"))
    page.mouse.down()
    page.mouse.move(*header(g, "cols", "D"), steps=5)
    page.mouse.up()
    g.wait_idle(0.2)
    assert editor(g) == "=SUM(C:D", editor(g)
    page.keyboard.type(")")
    page.keyboard.press("Enter")
    g.wait_text("A1", "110")
    g.click("A6")
    page.keyboard.type("=SUM(")
    page.mouse.click(*header(g, "rows", "2"))
    g.wait_idle(0.2)
    assert editor(g) == "=SUM(2:2", editor(g)
    page.keyboard.type(")")
    page.keyboard.press("Enter")
    g.wait_text("A6", "22")

    # GridCraft #845: Enter pastes the copied cells and ends copy mode.
    g.drag("C1", "C2")
    g.key("Mod+c")
    g.click("F1")
    page.keyboard.press("Enter")
    g.wait_text("F2", "2")
    assert g.text("F1") == "1"
    g.click("H1")
    page.keyboard.press("Enter")
    g.wait_idle(0.4)
    assert g.text("H1") in (None, "") and g.name_box() == "H2", (g.text("H1"), g.name_box())

    # Typing right after Enter in the name box goes into the cell it went to (the name box's own
    # text selection used to take the keyboard back).
    g.select("M300")
    page.keyboard.type("end")
    page.keyboard.press("Enter")
    g.select("M1")
    page.keyboard.type("top")
    page.keyboard.press("Enter")
    g.wait_text("M1", "top")
    assert g.name_box() == "M2", g.name_box()

    # GridCraft #762: Ctrl+Shift+Down scrolls to the end it extends to, not the anchor.
    g.select("M1")
    page.focus("#rg-editor")
    g.key(f"{MOD}+Shift+ArrowDown")
    page.wait_for_function("() => [...document.querySelectorAll('.rg-rowhead')].some((e) => e.innerText.trim() === '300')", timeout=8000)
    g.cell_rect("M300")

    # GridCraft #687: a selection down to the last row shows there, outline and all. Browsers stop
    # scrolling near 15 million px (row 750,000 at 100%); the rows past that are squeezed in.
    g.select("K5")
    page.focus("#rg-editor")
    g.key(f"{MOD}+Shift+ArrowDown")
    page.wait_for_function("() => [...document.querySelectorAll('.rg-rowhead')].some((e) => e.innerText.trim() === '1048576')", timeout=8000)
    g.wait_idle(0.5)
    sel = page.locator(".rg-selection").bounding_box()
    view = page.locator(".rg-scroller").bounding_box()
    assert view["y"] < sel["y"] + sel["height"] <= view["y"] + view["height"] + 2, (sel, view)
    g.select("D900000")
    page.keyboard.type("far")
    page.keyboard.press("Enter")
    g.wait_text("D900000", "far")
    g.key(f"{MOD}+Home")
    page.wait_for_function("() => [...document.querySelectorAll('.rg-rowhead')].some((e) => e.innerText.trim() === '1')", timeout=8000)
    g.select("A1")
    g.wait_idle(0.5)

    # GridCraft #668: a chart just inserted is selected, so Delete removes it, not its data.
    g.select("C1:D4")
    page.click("#tab-insert")
    page.click('button[title="Column"]')
    page.locator("[role=menuitem]").first.click()
    page.locator("[role=menu]").wait_for(state="detached")
    chart = page.locator(".rg-chart")
    chart.wait_for(state="visible")
    page.locator("#tab-chart-design").wait_for(state="visible")
    page.focus("#rg-editor")
    page.keyboard.press("Delete")
    chart.wait_for(state="detached")
    g.wait_text("C1", "1")
    g.wait_text("D4", "40")

    # GridCraft #507: double-clicking the fill handle fills down as far as the data beside it.
    g.type_in("E1", "=D1+1")
    g.click("E1")
    x, y, w, h = g.cell_rect("E1")
    page.mouse.dblclick(x + w - 1, y + h - 1)
    g.wait_text("E4", "41")
    assert g.text("E5") in (None, ""), g.text("E5")

    # GridCraft #777: a link can't run script; Ctrl/Cmd+click refuses it.
    page.evaluate("() => { window.__opened = null; window.open = (u) => { window.__opened = u; return null; }; }")
    g.select("A10")
    page.click("#tab-insert")
    page.get_by_role("button", name="Link", exact=True).click()
    page.locator("#dlg-hyperlink").wait_for(state="visible")
    page.locator("#dlg-hyperlink input").first.fill("Click me")
    page.fill("#dlg-target", "javascript:alert(document.cookie)//x.y")
    g.dialog_ok("hyperlink")
    g.wait_text("A10", "Click me")
    x, y, w, h = g.cell_rect("A10")
    page.keyboard.down(MOD)
    page.mouse.click(x + 12, y + h / 2)
    page.keyboard.up(MOD)
    toast(g, "opens only web")
    assert page.evaluate("() => window.__opened") is None
    g.shot("grid")

    # GridCraft #505: a list's input message hides while its list is open (and #479, #375 below).
    src = g.out / "validation.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    dv = DataValidation(type="list", formula1='"North,South,East"', showInputMessage=True, promptTitle="Region", prompt="Pick a region")
    ws.add_data_validation(dv)
    dv.add("B2")
    ws["D5"] = "Quarterly report"
    for c in "DEFG":
        ws[f"{c}5"].alignment = Alignment(horizontal="centerContinuous")
    ws["D7"] = 1234567.891
    ws["D7"].number_format = "#,##0.00"
    ws["D7"].alignment = Alignment(shrink_to_fit=True)
    ws["D8"] = 1234567.891
    ws["D8"].number_format = "#,##0.00"
    ws.column_dimensions["D"].width = 8
    ws["F10"] = "merged"
    ws.merge_cells("F10:H11")
    red = Side(style="medium", color="FF0000")
    for r in (10, 11):
        for c in "FGH":
            ws[f"{c}{r}"].border = Border(top=red if r == 10 else None, bottom=red if r == 11 else None, left=red if c == "F" else None, right=red if c == "H" else None)
    wb.save(src)
    page.click("#tab-file")
    page.click("#bs-open")
    page.locator("#dlg-open").wait_for(state="visible")
    page.locator('#dlg-open input[type="file"]').set_input_files(str(src))
    g.wait_title("validation")
    g.wait_idle(0.8)
    g.click("B2")
    msg = page.locator(".rg-input-msg")
    msg.wait_for(state="visible")
    page.locator(".rg-list-btn").click()
    page.locator(".rg-list").wait_for(state="visible")
    msg.wait_for(state="detached")
    page.keyboard.press("Escape")
    page.locator(".rg-list").wait_for(state="detached")
    msg.wait_for(state="visible")

    # GridCraft #479: Center Across Selection centres over the empty cells to its right that have it.
    d5 = page.locator('.rg-pane [data-cell="D5"]').bounding_box()
    x, y, w, h = g.cell_rect("D5")
    gx, gy, gw, gh = g.cell_rect("G5")
    assert abs(d5["x"] - x) < 2 and abs(d5["x"] + d5["width"] - (gx + gw)) < 3, (d5, x, gx + gw)
    # GridCraft #375: Shrink to Fit draws a smaller font where the number would show ###.
    assert g.text("D7") == "1,234,567.89" and g.text("D8").startswith("###"), (g.text("D7"), g.text("D8"))
    size = page.evaluate("() => parseFloat(getComputedStyle(document.querySelector('.rg-pane [data-cell=\"D7\"]')).fontSize)")
    assert size < 14, size
    # GridCraft #376: a merged area's right and bottom borders, kept on its right and bottom cells.
    g.click("B14")
    canvas = page.locator(".rg-view canvas").bounding_box()

    def pixel(x, y):
        return page.evaluate(
            "([x, y]) => { const cv = document.querySelector('.rg-view canvas'); const r = cv.getBoundingClientRect(); const k = cv.width / r.width;"
            " return [...cv.getContext('2d').getImageData(Math.round(x * k), Math.round(y * k), 1, 1).data].slice(0, 3); }",
            [x - canvas["x"], y - canvas["y"]],
        )

    fx, fy, fw, fh = g.cell_rect("F10")
    hx, hy, hw, hh = g.cell_rect("H11")
    mid_x, mid_y = (fx + hx + hw) / 2, (fy + hy + hh) / 2
    for side, (px_, py_) in {"left": (fx - 1, mid_y), "top": (mid_x, fy - 1), "right": (hx + hw - 1, mid_y), "bottom": (mid_x, hy + hh - 1)}.items():
        assert pixel(px_, py_) == [255, 0, 0], (side, pixel(px_, py_))

    # GridCraft #1907: a finger dragging on the grid pans it; it doesn't select cells.
    ctx = page.context.browser.new_context(viewport={"width": 900, "height": 700}, has_touch=True, is_mobile=False)
    try:
        touch = ctx.new_page()
        t = Grid(touch, g.out)
        t.open(g.page.url.split("/book/")[0])
        assert t.name_box() == "A1"
        box = touch.locator(".rg-scroller").bounding_box()
        # The finger's own events (Input.dispatchTouchEvent), not Input.synthesizeScrollGesture: in headless
        # Chromium on Linux (CI) the synthesized gesture's moves never reach the page as touchmove, so nothing
        # scrolls; the same gesture scrolls on macOS.
        cdp = ctx.new_cdp_session(touch)
        x, y = box["x"] + 300, box["y"] + 400
        cdp.send("Input.dispatchTouchEvent", {"type": "touchStart", "touchPoints": [{"x": x, "y": y}]})
        for i in range(1, 16):
            cdp.send("Input.dispatchTouchEvent", {"type": "touchMove", "touchPoints": [{"x": x, "y": y - 20 * i}]})
        cdp.send("Input.dispatchTouchEvent", {"type": "touchEnd", "touchPoints": []})
        touch.wait_for_function("() => document.querySelector('.rg-scroller').scrollTop > 100", timeout=8000)
        t.wait_idle(0.5)
        assert t.name_box() == "A1", t.name_box()
    finally:
        ctx.close()


if __name__ == "__main__":
    sys.exit(run("gridcraft_issues", body))
