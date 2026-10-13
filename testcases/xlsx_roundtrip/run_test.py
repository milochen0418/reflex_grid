"""Open an .xlsx made by openpyxl through File › Open, check what the grid shows, save it again
and read the downloaded file back with openpyxl (formulas, cached values, styles, merges)."""

import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from _lib.grid import Grid, run  # noqa: E402


def make_book(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(["Item", "Qty", "Price", "Total"])
    for item, q, p in (("Pen", 3, 1.5), ("Ink", 5, 4.25), ("Pad", 2, 3.0)):
        r = ws.max_row + 1
        ws.append([item, q, p, f"=B{r}*C{r}"])
    ws["D5"] = "=SUM(D2:D4)"
    ws["A1"].font = Font(bold=True)
    ws["D5"].number_format = '"$"#,##0.00'
    ws["A7"] = "Merged note"
    ws.merge_cells("A7:C7")
    ws["A8"] = "=TEXTJOIN(\", \",TRUE,A2:A4)"
    ws.column_dimensions["A"].width = 18
    ws["B1"].fill = PatternFill("solid", fgColor="FFFF00")
    s2 = wb.create_sheet("Second")
    s2["A1"] = "=Report!D5*2"
    wb.save(path)


def body(g: Grid) -> None:
    page = g.page
    src = g.out / "source.xlsx"
    make_book(src)
    page.click("#tab-file")
    page.click("#bs-open")
    page.locator("#dlg-open").wait_for(state="visible")
    page.locator('#dlg-open input[type="file"]').set_input_files(str(src))
    g.wait_title("source")
    g.wait_idle(1.0)
    # openpyxl writes no cached values: everything is calculated on open.
    g.wait_text("D2", "4.5")
    g.wait_text("D5", "$31.75")
    g.wait_text("A8", "Pen, Ink, Pad")
    g.click("B7")
    assert g.name_box() == "A7", g.name_box()
    page.locator(".rg-sheet-tab", has_text="Second").click()
    g.wait_text("A1", "63.5")
    page.locator(".rg-sheet-tab", has_text="Report").click()
    g.wait_idle(0.5)
    # Edit, then save (download) and read the result with openpyxl.
    g.type_in("B2", "10")
    g.wait_text("D5", "$42.25")
    with page.expect_download(timeout=20000) as dl:
        page.click("#qa-save")
    out = g.out / "saved.xlsx"
    dl.value.save_as(out)
    wb = openpyxl.load_workbook(out)
    ws = wb["Report"]
    assert ws["D2"].value == "=B2*C2", ws["D2"].value
    assert ws["D5"].value == "=SUM(D2:D4)", ws["D5"].value
    assert ws["A1"].font.bold, "bold lost"
    assert "A7:C7" in [str(r) for r in ws.merged_cells.ranges], ws.merged_cells.ranges
    assert ws["A8"].value.startswith("=_xlfn.TEXTJOIN("), ws["A8"].value
    vals = openpyxl.load_workbook(out, data_only=True)
    assert vals["Report"]["D5"].value == 42.25, vals["Report"]["D5"].value
    assert vals["Second"]["A1"].value == 84.5, vals["Second"]["A1"].value
    assert vals["Report"]["A8"].value == "Pen, Ink, Pad"
    g.shot("roundtrip")


if __name__ == "__main__":
    sys.exit(run("xlsx_roundtrip", body))
